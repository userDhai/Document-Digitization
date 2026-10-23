"""Privacy-first document processing pipeline for the prototype."""
from __future__ import annotations

import hashlib
import io
import json
import math
import os
import re
import shutil
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import cv2
import numpy as np
import pytesseract
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

DEFAULT_TESSERACT_CMD = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
DEFAULT_OCR_LANG = "eng+ara"
PDF_RENDER_DPI = max(150, min(600, int(os.getenv("PDF_RENDER_DPI", "300"))))

BASE_DIR = Path(__file__).resolve().parents[1]
TEMP_DIR = BASE_DIR / "temporary"
OUTPUT_DIR = BASE_DIR / "outputs"
TOKEN_RE = re.compile(r"\[\[([A-Z]+(?:_[A-Z]+)*_\d{3})\]\]")
PATTERNS = {
    "EMAIL": re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I),
    "PHONE": re.compile(r"(?<!\w)(?:\+?\d[\d .()\-]{7,}\d)(?!\w)"),
    "SSN": re.compile(r"\b\d{3}-\d{2}-\d{4}\b"),
    "CARD": re.compile(r"\b(?:\d[ -]?){13,19}\b"),
    "DATE_OF_BIRTH": re.compile(r"\b(?:DOB|date of birth)\s*:?\s*\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b", re.I),
}

# Context labels make names and identifiers detectable even when the value itself
# has no recognizable format. Arabic labels are escaped to keep this file portable.
_NAME_LABELS = (
    r"full\s+name|first\s+name|given\s+name|middle\s+name|last\s+name|family\s+name|surname|"
    r"recipient(?:\s+name)?|holder(?:\s+name)?|employee\s+name|"
    r"customer\s+name|applicant(?:\s+name)?|patient(?:\s+name)?|name|"
    r"\u0627\u0644\u0627\u0633\u0645\s+(?:\u0627\u0644\u0623\u0648\u0644|\u0627\u0644\u0623\u0648\u0633\u0637|\u0627\u0644\u0623\u062e\u064a\u0631)|"
    r"\u0627\u0633\u0645\s+\u0627\u0644\u0639\u0627\u0626\u0644\u0629|\u0627\u0644\u0644\u0642\u0628|"
    r"\u0627\u0644\u0627\u0633\u0645|"
    r"\u0627\u0633\u0645\s+\u0627\u0644\u0645\u0633\u062a\u0644\u0645|"
    r"\u0627\u0633\u0645\s+\u0627\u0644\u0645\u0648\u0638\u0641|"
    r"\u0627\u0633\u0645\s+\u062d\u0627\u0645\u0644|"
    r"\u0627\u0644\u0645\u0633\u062a\u0644\u0645|"
    r"\u0645\u0642\u062f\u0645\s+\u0627\u0644\u0637\u0644\u0628"
)
_FIELD_LABELS = re.compile(
    r"\s+(?:first\s+name|given\s+name|middle\s+name|last\s+name|family\s+name|surname|"
    r"\u0627\u0644\u0627\u0633\u0645\s+(?:\u0627\u0644\u0623\u0648\u0644|\u0627\u0644\u0623\u0648\u0633\u0637|\u0627\u0644\u0623\u062e\u064a\u0631)|"
    r"\u0627\u0633\u0645\s+\u0627\u0644\u0639\u0627\u0626\u0644\u0629|\u0627\u0644\u0644\u0642\u0628|"
    r"employee\s*(?:id|number|no\.?)|national\s*(?:id|identity)|identity\s*(?:number|no\.?)|"
    r"passport(?:\s*(?:number|no\.?))?|id|phone|mobile|email|department|city|state|zip|postal\s+code|"
    r"start\s+date|date|address|date\s+of\s+birth|dob|gender|title|salary|"
    r"designation|position|nationality|passport|company|signature)\s*[:#-]", re.I)
_PERSON_WORDS = re.compile(r"[\w\u0600-\u06ff][\w\u0600-\u06ff'.-]*", re.UNICODE)
_ID_LABEL = re.compile(r"\b(?:employee\s*(?:id|number|no\.?))\s*[:#-]\s*([A-Z0-9][A-Z0-9-]{4,})\b", re.I)
_GENERIC_ID_LABEL = re.compile(r"\bID\s*[:#-]\s*([A-Z0-9][A-Z0-9-]{4,})\b", re.I)
_NATIONAL_ID_LABEL = re.compile(
    r"\b(?:national\s*(?:id|identity)|identity\s*(?:number|no\.?))\s*[:#-]\s*([A-Z0-9][A-Z0-9-]{4,})\b", re.I)
_PASSPORT_LABEL = re.compile(r"\bpassport(?:\s*(?:number|no\.?))?\s*[:#-]\s*([A-Z0-9][A-Z0-9-]{4,})\b", re.I)
_ADDRESS_LABELS = (
    r"home\s+address|street\s+address|mailing\s+address|residential\s+address|"
    r"billing\s+address|address|\u0627\u0644\u0639\u0646\u0648\u0627\u0646|"
    r"\u0639\u0646\u0648\u0627\u0646\s+\u0627\u0644\u0633\u0643\u0646"
)


@dataclass
class Page:
    image: np.ndarray
    text: str
    words: list[dict[str, Any]]
    number: int = 0
    ocr_diagnostics: dict[str, Any] | None = None


class GeminiConfigurationError(RuntimeError):
    pass


class OCRConfigurationError(ValueError):
    pass


class GeminiRequestError(RuntimeError):
    def __init__(self, exception_type: str, status_code: Any = None, status: Any = None):
        # Provider exception messages can echo request data, so retain safe fields only.
        self.exception_type = exception_type
        self.status_code = status_code if isinstance(status_code, (int, str)) else None
        self.status = status if isinstance(status, str) else None
        super().__init__("Gemini request failed")


class GeminiResponseError(ValueError):
    pass


def load_document(path: Path, ocr_cache: list[dict[str, Any]] | None = None) -> list[Page]:
    """Rasterize PDF pages or load a supported image; OCR locally before any API call."""
    cached_by_page = {item["page"]: item for item in (ocr_cache or [])}
    if not ocr_cache:
        # Resolve at processing time so values loaded from .env are honored.
        pytesseract.pytesseract.tesseract_cmd = os.getenv("TESSERACT_CMD", DEFAULT_TESSERACT_CMD)
        ocr_lang = os.getenv("TESSERACT_LANG", DEFAULT_OCR_LANG)
        required_languages = {language for language in ("eng", "ara") if language in ocr_lang.split("+")}
        if required_languages != {"eng", "ara"}:
            raise OCRConfigurationError("TESSERACT_LANG must include both eng and ara for English and Arabic recognition.")
        try:
            installed_languages = set(pytesseract.get_languages(config=""))
        except Exception as exc:
            raise OCRConfigurationError("Tesseract could not be started. Check TESSERACT_CMD and install its language data.") from exc
        missing = required_languages - installed_languages
        if missing:
            missing_names = ", ".join("Arabic (ara)" if item == "ara" else "English (eng)" for item in sorted(missing))
            raise OCRConfigurationError(f"Tesseract language data is missing: {missing_names}. Install both eng and ara traineddata files.")
    else:
        ocr_lang = os.getenv("TESSERACT_LANG", DEFAULT_OCR_LANG)
    images: list[np.ndarray] = []
    if path.suffix.lower() == ".pdf":
        import fitz
        doc = fitz.open(path)
        try:
            for page in doc:
                pix = page.get_pixmap(dpi=PDF_RENDER_DPI, alpha=False)
                images.append(cv2.cvtColor(np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, pix.n), cv2.COLOR_RGB2BGR))
        finally:
            doc.close()
    else:
        with Image.open(path) as source:
            rgb = np.asarray(source.convert("RGB"))
            images.append(cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR))
    pages = []
    for page_number, image in enumerate(images):
        cached = cached_by_page.get(page_number)
        if cached:
            pages.append(Page(image, cached["text"], [dict(word) for word in cached["words"]], page_number,
                              dict(cached.get("ocr_diagnostics") or {})))
            continue
        data = pytesseract.image_to_data(image, lang=ocr_lang, config="--oem 1 --psm 3",
                                         output_type=pytesseract.Output.DICT)
        words = []
        for i, value in enumerate(data["text"]):
            value = value.strip()
            if value:
                words.append({"text": value, "x": int(data["left"][i]), "y": int(data["top"][i]),
                              "w": int(data["width"][i]), "h": int(data["height"][i]),
                              "block": int(data["block_num"][i]),
                              "paragraph": int(data.get("par_num", [0] * len(data["text"]))[i]),
                              "line": int(data["line_num"][i]), "word": int(data.get("word_num", list(range(len(data["text"]))))[i]),
                              "confidence": float(data.get("conf", [-1] * len(data["text"]))[i])})
        # Retain Tesseract's block/line order and separate lines so unrelated
        # columns and form fields do not become one ambiguous OCR sentence.
        line_words: dict[tuple[int, int, int], list[dict[str, Any]]] = {}
        for word in words:
            line_words.setdefault((word["block"], word["paragraph"], word["line"]), []).append(word)
        words = [word for line in line_words.values() for word in line]
        text = "\n".join(" ".join(item["text"] for item in line) for line in line_words.values())
        confidences = [word["confidence"] for word in words if word["confidence"] >= 0]
        diagnostics = {"languages": ocr_lang, "pdf_render_dpi": PDF_RENDER_DPI if path.suffix.lower() == ".pdf" else None,
                       "page_width": int(image.shape[1]), "page_height": int(image.shape[0]),
                       "word_count": len(words), "line_count": len(line_words),
                       "average_confidence": round(sum(confidences) / len(confidences), 1) if confidences else None}
        pages.append(Page(image, text, words, page_number, diagnostics))
    return pages


def _sensitive_spans(text: str) -> list[tuple[int, int, str, str]]:
    spans = []
    for kind, pattern in PATTERNS.items():
        for match in pattern.finditer(text):
            if kind == "PHONE" and re.search(
                    r"(?:employee\s*(?:id|number|no\.?)|national\s*(?:id|identity)|"
                    r"identity\s*(?:number|no\.?)|passport(?:\s*(?:number|no\.?))?|id)\s*[:#-]\s*$",
                    text[max(0, match.start() - 48):match.start()], re.I):
                continue
            value = match.group(0)
            if kind == "PHONE" and len(re.sub(r"\D", "", value)) < 9:
                continue
            if kind == "CARD":
                digits = re.sub(r"\D", "", value)
                if not 13 <= len(digits) <= 19:
                    continue
            spans.append((match.start(), match.end(), kind, value))
    # Extract a name only when it follows an explicit person/recipient label.
    for label in re.finditer(rf"(?i)(?<!\w)(?:{_NAME_LABELS})\s*[:#-]\s*", text):
        start = label.end()
        stop = re.search(r"[,;|\n]|\s{2,}", text[start:])
        end = start + stop.start() if stop else len(text)
        next_field = _FIELD_LABELS.search(text, start, end)
        if next_field:
            end = next_field.start()
        candidate = text[start:end].strip()
        if TOKEN_RE.fullmatch(candidate):
            continue
        words = list(_PERSON_WORDS.finditer(candidate))
        if not words:
            continue
        # Keep at most five name-like words, excluding trailing form text.
        candidate_end = words[min(len(words), 5) - 1].end()
        name = candidate[:candidate_end].strip(" .:-")
        if name:
            name_start = start + len(text[start:end]) - len(text[start:end].lstrip())
            spans.append((name_start, name_start + len(name), "PERSON", name))
    # Capture labelled employee numbers, including values that are not phone-shaped.
    for match in _ID_LABEL.finditer(text):
        spans.append((match.start(1), match.end(1), "EMPLOYEE_ID", match.group(1)))
    for match in _GENERIC_ID_LABEL.finditer(text):
        if re.search(r"national\s*$", text[max(0, match.start() - 40):match.start()], re.I):
            continue
        spans.append((match.start(1), match.end(1), "ID", match.group(1)))
    for pattern, category in ((_NATIONAL_ID_LABEL, "NATIONAL_ID"), (_PASSPORT_LABEL, "PASSPORT")):
        for match in pattern.finditer(text):
            spans.append((match.start(1), match.end(1), category, match.group(1)))
    # Address values have no reliable standalone format, so detect them by field label.
    for label in re.finditer(rf"(?i)(?<!\w)(?:{_ADDRESS_LABELS})\s*[:#-]\s*", text):
        start = label.end()
        stop = re.search(r"[;|\n]|\s{2,}", text[start:])
        end = start + stop.start() if stop else len(text)
        next_field = _FIELD_LABELS.search(text, start, end)
        if next_field:
            end = next_field.start()
        address = text[start:end].strip(" ,;:-")
        if address and not TOKEN_RE.fullmatch(address):
            leading = len(text[start:end]) - len(text[start:end].lstrip())
            spans.append((start + leading, start + leading + len(address), "ADDRESS", address))
    spans.sort(key=lambda x: (x[0], -(x[1] - x[0])))
    selected = []
    for span in spans:
        if not selected or span[0] >= selected[-1][1]:
            selected.append(span)
    return selected


def detect_and_tokenize(text: str, extra_spans: list[tuple[int, int, str, str]] | None = None
                        ) -> tuple[str, dict[str, str], list[dict[str, str]]]:
    spans = [*_sensitive_spans(text), *(extra_spans or [])]
    spans.sort(key=lambda span: (span[0], -(span[1] - span[0])))
    selected_spans = []
    for span in spans:
        if 0 <= span[0] < span[1] <= len(text) and (not selected_spans or span[0] >= selected_spans[-1][1]):
            selected_spans.append(span)
    counts: dict[str, int] = {}
    mapping, findings, out, cursor = {}, [], [], 0
    for start, end, kind, value in selected_spans:
        counts[kind] = counts.get(kind, 0) + 1
        token = f"[[{kind}_{counts[kind]:03d}]]"
        out.extend((text[cursor:start], token))
        cursor = end
        mapping[token] = value
        findings.append({"token": token, "category": kind, "masked_value": "Detected and tokenized locally"})
    out.append(text[cursor:])
    return "".join(out), mapping, findings


def _tokenize_spans(text: str, spans: list[tuple[int, int, str, str]]
                    ) -> tuple[str, dict[str, str], list[dict[str, str]], list[tuple[int, int, str]]]:
    """Tokenize only the user-selected local spans; detection itself always runs independently."""
    selected = []
    for span in sorted(spans, key=lambda item: (item[0], -(item[1] - item[0]))):
        if 0 <= span[0] < span[1] <= len(text) and (not selected or span[0] >= selected[-1][1]):
            selected.append(span)
    counts: dict[str, int] = {}
    mapping, findings, out, cursor, tokens = {}, [], [], 0, []
    for start, end, kind, value in selected:
        counts[kind] = counts.get(kind, 0) + 1
        token = f"[[{kind}_{counts[kind]:03d}]]"
        out.extend((text[cursor:start], token)); cursor = end
        mapping[token] = value
        findings.append({"token": token, "category": kind, "masked_value": "Detected and tokenized locally"})
        tokens.append((start, end, token))
    out.append(text[cursor:])
    return "".join(out), mapping, findings, tokens


def _ocr_review_items(pages: list[Page], submitted: list[dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    """Group OCR words by text line while retaining original page coordinates."""
    rows = []
    for page in pages:
        lines: dict[tuple[int, int, int], list[dict[str, Any]]] = {}
        for word in page.words:
            lines.setdefault((word.get("block", 0), word.get("paragraph", 1), word.get("line", 0)), []).append(word)
        for (block, paragraph, line), words in lines.items():
            x0, y0 = min(w["x"] for w in words), min(w["y"] for w in words)
            x1 = max(w["x"] + w["w"] for w in words)
            y1 = max(w["y"] + w["h"] for w in words)
            text = " ".join(w["text"] for w in words)
            base_id = f"ocr-p{page.number}-b{block}-l{line}"
            row_id = (f"ocr-p{page.number}-b{block}-q{paragraph}-l{line}"
                      if any(row["id"] == base_id for row in rows)
                      else base_id)
            matches = _sensitive_spans(text)
            rows.append({"id": row_id, "page": page.number, "text": text, "original_text": text,
                         "category": ", ".join(dict.fromkeys(item[2] for item in matches)),
                         "detection_status": "sensitive" if matches else "non-sensitive",
                         "classification": "sensitive" if matches else "non-sensitive",
                         "transmission": "local" if matches else "include",
                         "manual_sensitive": False, "x": x0, "y": y0,
                         "width": x1 - x0, "height": y1 - y0,
                         "block": block, "paragraph": paragraph, "line": line,
                         "mask_id": None, "token": None, "is_manual": False})
        if not lines and page.text.strip():
            row_id = f"ocr-p{page.number}-text"
            matches = _sensitive_spans(page.text)
            rows.append({"id": row_id, "page": page.number, "text": page.text,
                         "original_text": page.text,
                         "category": ", ".join(dict.fromkeys(item[2] for item in matches)),
                         "detection_status": "sensitive" if matches else "non-sensitive",
                         "classification": "sensitive" if matches else "non-sensitive",
                         "transmission": "local" if matches else "include",
                         "manual_sensitive": False, "x": None, "y": None,
                         "width": None, "height": None, "mask_id": None,
                         "block": None, "paragraph": None, "line": None,
                         "token": None, "is_manual": True})
    overrides = {item.get("id"): item for item in (submitted or []) if isinstance(item, dict)}
    by_id = {row["id"]: row for row in rows}
    for row_id, override in overrides.items():
        if row_id in by_id:
            row = by_id[row_id]
            text = override.get("text", row["text"])
            if not isinstance(text, str) or len(text) > 12000:
                raise ValueError("OCR text edits must be text under 12,000 characters.")
            row["text"] = text
            row["classification"] = override.get("classification", row["classification"])
            row["transmission"] = override.get("transmission", row["transmission"])
            row["manual_sensitive"] = bool(override.get("manual_sensitive", False))
            row["manual_mask_id"] = override.get("manual_mask_id")
        elif override.get("is_manual") is True:
            text = override.get("text", "")
            page = override.get("page")
            if (not isinstance(text, str) or not text.strip() or len(text) > 12000 or
                    isinstance(page, bool) or not isinstance(page, int) or page not in {p.number for p in pages}):
                raise ValueError("Manually added OCR text must include text and a valid page.")
            rows.append({"id": str(row_id), "page": page, "text": text, "original_text": "",
                         "category": "Manual review", "detection_status": "uncertain",
                         "classification": override.get("classification", "uncertain"),
                         "transmission": override.get("transmission", "local"),
                         "manual_sensitive": bool(override.get("manual_sensitive", True)),
                         "x": None, "y": None, "width": None, "height": None,
                         "block": None, "line": None,
                         "mask_id": override.get("mask_id"), "token": None, "is_manual": True})
        else:
            raise ValueError("An OCR review item is invalid.")
    for row in rows:
        if row["classification"] not in {"sensitive", "non-sensitive", "uncertain"}:
            raise ValueError("Choose sensitive, non-sensitive, or uncertain for OCR items.")
        if row["transmission"] not in {"local", "include"}:
            raise ValueError("Choose whether each OCR item stays local or is included for AI.")
        matches = _sensitive_spans(row["text"])
        row["category"] = ", ".join(dict.fromkeys(item[2] for item in matches)) or (
            "Manually marked" if row["manual_sensitive"] else "")
        row["detection_status"] = "sensitive" if matches or row["manual_sensitive"] or row["classification"] == "sensitive" else (
            "uncertain" if row["classification"] == "uncertain" else "non-sensitive")
    return rows


def _automatic_masks(pages: list[Page], mappings: dict[str, str]) -> list[dict[str, Any]]:
    """Return safe, editable rectangle suggestions for every detected value occurrence."""
    masks = []
    for token, value in mappings.items():
        for page in pages:
            needle, lower = value.casefold(), page.text.casefold()
            offset = 0
            occurrence = 0
            while (position := lower.find(needle, offset)) >= 0:
                end = position + len(needle)
                char_position = 0
                selected = []
                for word in page.words:
                    begin, finish = char_position, char_position + len(word["text"])
                    if finish > position and begin < end:
                        selected.append(word)
                    char_position = finish + 1
                if selected:
                    occurrence += 1
                    left = max(0, min(word["x"] for word in selected) - 3)
                    top = max(0, min(word["y"] for word in selected) - 3)
                    right = min(page.image.shape[1], max(word["x"] + word["w"] for word in selected) + 3)
                    bottom = min(page.image.shape[0], max(word["y"] + word["h"] for word in selected) + 3)
                    masks.append({"id": f"auto-{token[2:-2]}-p{page.number}-n{occurrence}",
                                  "page": page.number, "x": left, "y": top,
                                  "width": right - left, "height": bottom - top,
                                  "kind": "automatic", "category": token[2:-6]})
                offset = end
    return masks


def _redact_visuals(pages: list[Page], mappings: dict[str, str],
                    automatic_masks: list[dict[str, Any]],
                    manual_masks: list[dict[str, int]] | None = None) -> tuple[list[str], bool]:
    """Render the selected masks and verify every known sensitive OCR word is blacked out."""
    previews = []
    all_covered = True
    safe_dir = TEMP_DIR / uuid.uuid4().hex
    safe_dir.mkdir(parents=True, exist_ok=True)
    page_numbers = [page.number for page in pages]
    for index, page in enumerate(pages, 1):
        img = page.image.copy()
        lower = page.text.lower()
        for mask in automatic_masks:
            if mask["page"] != page.number:
                continue
            x, y = mask["x"], mask["y"]
            x1, y1 = x + mask["width"], y + mask["height"]
            if x < 0 or y < 0 or x1 > img.shape[1] or y1 > img.shape[0]:
                all_covered = False
                continue
            cv2.rectangle(img, (x, y), (x1, y1), (0, 0, 0), -1)
            if mask.get("required"):
                all_covered = all_covered and bool(np.all(img[y:y1, x:x1] == 0))
        for mask in manual_masks or []:
            if mask["page"] != page.number:
                continue
            x, y = mask["x"], mask["y"]
            x1, y1 = x + mask["width"], y + mask["height"]
            if x < 0 or y < 0 or x1 > img.shape[1] or y1 > img.shape[0]:
                all_covered = False
                continue
            cv2.rectangle(img, (x, y), (x1, y1), (0, 0, 0), -1)
            all_covered = all_covered and bool(np.all(img[y:y1, x:x1] == 0))
        # Verify coverage after applying both automatic and manually adjusted masks.
        for value in mappings.values():
            needle = value.lower()
            offset = 0
            while (pos := lower.find(needle, offset)) >= 0:
                end = pos + len(needle)
                char_pos = 0
                selected = []
                for word in page.words:
                    begin, finish = char_pos, char_pos + len(word["text"])
                    if finish > pos and begin < end:
                        selected.append(word)
                    char_pos = finish + 1
                for word in selected:
                    x, y, w, h = word["x"], word["y"], word["w"], word["h"]
                    all_covered = all_covered and bool(np.all(img[y:y + h, x:x + w] == 0))
                offset = end
        path = safe_dir / f"redacted-{index}.png"
        if not cv2.imwrite(str(path), img):
            raise ValueError("Privacy verification failed. The document was not sent to external AI.")
        previews.append(f"/api/preview/{safe_dir.name}/{path.name}")
    # Verify that every unique detected value occurred in OCR text and has a covered box.
    combined = "\n".join(page.text for page in pages).casefold()
    all_covered = all_covered and all(value.casefold() in combined for value in mappings.values())
    return previews, all_covered


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def verify_privacy(outgoing_payload: str, mapping: dict[str, str], visual_verified: bool = True,
                   preview_paths: list[str] | None = None, preview_hashes: list[str] | None = None,
                   manual_masks: list[dict[str, int]] | None = None,
                   page_numbers: list[int] | None = None,
                   approved_transmissions: list[str] | None = None) -> bool:
    """Check tokenized text and integrity/coverage of the exact sanitized images."""
    safe = outgoing_payload.casefold()
    if not visual_verified:
        return False
    if any(not value or value.casefold() in safe for value in mapping.values()):
        return False
    if any(token not in outgoing_payload for token in mapping):
        return False
    approved = {value.casefold() for value in (approved_transmissions or [])}
    if any(value.casefold() not in approved for _, _, _, value in _sensitive_spans(outgoing_payload)):
        return False
    if preview_paths is not None:
        if not preview_paths or (preview_hashes is not None and len(preview_hashes) != len(preview_paths)):
            return False
        images = []
        for index, item in enumerate(preview_paths):
            path = Path(item)
            if not path.is_file():
                return False
            if preview_hashes is not None and _file_sha256(path) != preview_hashes[index]:
                return False
            image = cv2.imread(str(path))
            if image is None:
                return False
            images.append(image)
        page_to_image = {number: index for index, number in enumerate(page_numbers or range(len(images)))}
        for mask in manual_masks or []:
            page = page_to_image.get(mask["page"])
            if page is None:
                return False
            image = images[page]
            x, y = mask["x"], mask["y"]
            x1, y1 = x + mask["width"], y + mask["height"]
            if x < 0 or y < 0 or x1 > image.shape[1] or y1 > image.shape[0]:
                return False
            if not np.all(image[y:y1, x:x1] == 0):
                return False
    elif manual_masks:
        return False
    return True


def _validate_json(candidate: str) -> dict[str, Any]:
    candidate = (candidate or "").strip()
    if not candidate:
        raise GeminiResponseError("Gemini returned an empty response.")
    candidate = re.sub(r"^```(?:json)?\s*|\s*```$", "", candidate, flags=re.I).strip()
    try:
        value = json.loads(candidate)
    except json.JSONDecodeError:
        start, end = candidate.find("{"), candidate.rfind("}")
        if start < 0 or end <= start:
            raise GeminiResponseError("Gemini returned a response that could not be parsed.") from None
        try:
            value = json.loads(candidate[start:end + 1])
        except json.JSONDecodeError:
            raise GeminiResponseError("Gemini returned a response that could not be parsed.") from None
    if not isinstance(value, dict):
        raise GeminiResponseError("Gemini returned data that does not match the expected extraction format.")
    if not isinstance(value.get("document_type"), str) or not isinstance(value.get("summary"), str):
        raise GeminiResponseError("Gemini returned data that does not match the expected extraction format.")
    if not isinstance(value.get("fields"), (dict, list)):
        raise GeminiResponseError("Gemini returned data that does not match the expected extraction format.")
    return value


def build_gemini_payload(safe_text: str, output_language: str = "en") -> str:
    """Build the complete outgoing text payload before privacy verification."""
    if output_language not in {"en", "ar"}:
        raise ValueError("Output language must be en or ar.")
    instruction = ("Use Arabic only for field labels and other output labels. " if output_language == "ar"
                   else "Use English only for field labels and other output labels. ")
    return ("Return only JSON with document_type, summary, and fields as an object. "
            f"{instruction}Treat document_type, summary, and every fields value as document content, not labels: "
            "keep them in the source language and preserve source wording exactly. Do not translate, paraphrase, "
            "normalize, or infer names, addresses, descriptions, company names, identifiers, dates, numbers, or any "
            "other document values. Field labels are keys and may use the selected output language; their associated "
            "values must remain unchanged. Preserve privacy tokens exactly.\n\n" + safe_text)


def call_gemini(verified_payload: str, mapping: dict[str, str], visual_verified: bool,
                image_paths: list[str] | None = None, image_hashes: list[str] | None = None,
                manual_masks: list[dict[str, int]] | None = None,
                page_numbers: list[int] | None = None,
                approved_transmissions: list[str] | None = None,
                human_approved: bool = False) -> tuple[dict[str, Any], str]:
    """Send prepared OCR and redacted page images after the user's approval."""
    if not image_paths:
        raise ValueError("Privacy verification failed. The document was not sent to external AI.")
    if not human_approved and not verify_privacy(verified_payload, mapping, visual_verified, image_paths, image_hashes,
                                                 manual_masks, page_numbers, approved_transmissions):
        raise ValueError("Privacy verification failed. The document was not sent to external AI.")
    key = os.getenv("GEMINI_API_KEY")
    if os.getenv("DEMO_MODE", "false").lower() == "true":
        arabic = "Use Arabic only for field labels" in verified_payload
        return {"document_type": "Document",
                "summary": "Extracted from privacy-filtered source OCR.",
                "fields": {"النص المستخرج" if arabic else "extracted_text":
                           verified_payload.split("\n\n", 1)[-1][:1000]}}, "demo"
    if not key:
        raise GeminiConfigurationError("Gemini API key is not configured.")
    try:
        from google import genai
        from google.genai import types
    except ImportError:
        raise GeminiConfigurationError("Gemini SDK is not installed. Install project requirements.") from None
    images = []
    try:
        max_edge = int(os.getenv("GEMINI_MAX_IMAGE_EDGE", "1800"))
    except ValueError:
        max_edge = 1800
    max_edge = max(800, min(3000, max_edge))
    for index, image_path in enumerate(image_paths):
        image_bytes = Path(image_path).read_bytes()
        if image_hashes is None or index >= len(image_hashes) or hashlib.sha256(image_bytes).hexdigest() != image_hashes[index]:
            raise ValueError("Privacy verification failed. The document was not sent to external AI.")
        try:
            with Image.open(io.BytesIO(image_bytes)) as source:
                image = source.convert("RGB")
                image.thumbnail((max_edge, max_edge), Image.Resampling.LANCZOS)
                encoded = io.BytesIO()
                image.save(encoded, format="JPEG", quality=88)
                images.append(types.Part.from_bytes(data=encoded.getvalue(), mime_type="image/jpeg"))
        except Exception:
            raise ValueError("Privacy verification failed. The document was not sent to external AI.") from None
    try:
        client = genai.Client(api_key=key)
        thinking_level = os.getenv("GEMINI_THINKING_LEVEL", "low").lower()
        if thinking_level not in {"low", "medium", "high"}:
            thinking_level = "low"
        response = client.models.generate_content(
            model=os.getenv("GEMINI_MODEL", "gemini-3.8-flash"),
            contents=[verified_payload, *images],
            config={"response_mime_type": "application/json",
                    "thinking_config": {"thinking_level": thinking_level}},
        )
    except Exception as exc:
        raise GeminiRequestError(type(exc).__name__, getattr(exc, "code", None), getattr(exc, "status", None)) from None
    return _validate_json(response.text or ""), "gemini"


def restore_tokens(value: Any, mapping: dict[str, str]) -> Any:
    if isinstance(value, str):
        return TOKEN_RE.sub(lambda m: mapping.get(m.group(0), m.group(0)), value)
    if isinstance(value, list):
        return [restore_tokens(item, mapping) for item in value]
    if isinstance(value, dict):
        return {restore_tokens(key, mapping) if isinstance(key, str) else key: restore_tokens(item, mapping)
                for key, item in value.items()}
    return value


def _manual_mask_spans(pages: list[Page], masks: list[dict[str, Any]]) -> list[tuple[int, int, str, str]]:
    spans = []
    page_offset = 0
    for page in pages:
        for mask in (item for item in masks if item["page"] == page.number):
            x0, y0 = mask["x"], mask["y"]
            x1, y1 = x0 + mask["width"], y0 + mask["height"]
            char_position = 0
            selected = []
            for word in page.words:
                begin, end = char_position, char_position + len(word["text"])
                wx, wy, ww, wh = word["x"], word["y"], word["w"], word["h"]
                if wx < x1 and wx + ww > x0 and wy < y1 and wy + wh > y0:
                    selected.append((begin, end))
                char_position = end + 1
            if selected:
                start, end = selected[0][0], selected[-1][1]
                spans.append((page_offset + start, page_offset + end, "PERSON_MANUAL",
                              page.text[start:end]))
        page_offset += len(page.text) + 1
    return spans


def _normalize_manual_masks(pages: list[Page], masks: list[dict[str, Any]] | None
                            ) -> list[dict[str, int]]:
    page_lookup = {page.number: page for page in pages}
    normalized = []
    for mask in masks or []:
        if not isinstance(mask, dict):
            raise ValueError("Manual mask coordinates are invalid.")
        page, x, y, width, height = (mask.get(name) for name in ("page", "x", "y", "width", "height"))
        if isinstance(page, bool) or not isinstance(page, int) or page not in page_lookup:
            raise ValueError("Manual mask page is invalid.")
        coords = (x, y, width, height)
        if any(isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value)
               for value in coords):
            raise ValueError("Manual mask coordinates are invalid.")
        x, y, width, height = (round(value) for value in coords)
        image_height, image_width = page_lookup[page].image.shape[:2]
        if x < 0 or y < 0 or width <= 0 or height <= 0 or x + width > image_width or y + height > image_height:
            raise ValueError("Manual mask must stay within the document page.")
        normalized.append({"id": str(mask.get("id") or f"manual-{uuid.uuid4().hex}"),
                           "kind": "manual", "page": page, "x": x, "y": y,
                           "width": width, "height": height})
    return normalized


def _normalize_keep_pages(available_pages: list[Page], keep_pages: list[int] | None) -> list[Page]:
    lookup = {page.number: page for page in available_pages}
    selected_numbers = list(lookup) if keep_pages is None else keep_pages
    if (not isinstance(selected_numbers, list) or not selected_numbers or
            any(isinstance(number, bool) or not isinstance(number, int) or number not in lookup
                for number in selected_numbers) or len(set(selected_numbers)) != len(selected_numbers)):
        raise ValueError("Select at least one valid document page.")
    return [lookup[number] for number in selected_numbers]


def _prepare_auto_masks(pages: list[Page], mappings: dict[str, str],
                        disabled_ids: list[str] | None,
                        overrides: list[dict[str, Any]] | None) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    automatic = _automatic_masks(pages, mappings)
    by_id = {mask["id"]: mask for mask in automatic}
    for override in overrides or []:
        if not isinstance(override, dict) or override.get("id") not in by_id:
            raise ValueError("An automatic mask selection is invalid.")
        original = by_id[override["id"]]
        revised = _normalize_manual_masks(pages, [{**override, "page": original["page"]}])[0]
        original.update({key: revised[key] for key in ("x", "y", "width", "height")})
    disabled = disabled_ids or []
    if any(not isinstance(mask_id, str) or mask_id not in by_id for mask_id in disabled):
        raise ValueError("An automatic mask selection is invalid.")
    for mask in automatic:
        mask["enabled"] = mask["id"] not in disabled
    selected = [mask for mask in automatic if mask["enabled"]]
    return automatic, selected


def prepare_document(path: Path, manual_masks: list[dict[str, Any]] | None = None,
                     keep_pages: list[int] | None = None,
                     disabled_automatic_mask_ids: list[str] | None = None,
                     automatic_mask_overrides: list[dict[str, Any]] | None = None,
                     ocr_review: list[dict[str, Any]] | None = None,
                     include_confirmed: bool = False,
                     output_language: str = "en",
                     ocr_cache: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    """Rebuild masks and privacy state; reuse server-local OCR for unchanged source pages."""
    TEMP_DIR.mkdir(exist_ok=True)
    source_pages = load_document(path, ocr_cache)
    page_ocr_cache = [{"page": page.number, "text": page.text, "words": page.words,
                       "ocr_diagnostics": page.ocr_diagnostics} for page in source_pages]
    pages = _normalize_keep_pages(source_pages, keep_pages)
    masks = _normalize_manual_masks(pages, manual_masks)
    rows = _ocr_review_items(pages, ocr_review)
    unmasked_manual_rows = []
    for row in rows:
        if row["x"] is not None:
            covering = next((mask for mask in masks if mask["page"] == row["page"] and
                mask["x"] < row["x"] + row["width"] and row["x"] < mask["x"] + mask["width"] and
                mask["y"] < row["y"] + row["height"] and row["y"] < mask["y"] + mask["height"]), None)
            if covering:
                row["manual_sensitive"] = True
                row["classification"] = "sensitive"
                row["detection_status"] = "sensitive"
                row["category"] = row["category"] or "Manual mask"
                row["manual_mask_id"] = covering["id"]
        if row["is_manual"] and row["detection_status"] == "sensitive":
            linked = next((mask for mask in masks if mask["id"] == row.get("mask_id") and mask["page"] == row["page"]), None)
            if linked is None:
                unmasked_manual_rows.append(row["id"])
    raw_text = "\n".join(row["text"] for row in rows)
    offsets, offset = {}, 0
    for row in rows:
        offsets[row["id"]] = offset
        offset += len(row["text"]) + 1
    manual_spans: list[tuple[int, int, str, str]] = []
    for row in rows:
        if row["x"] is None:
            continue
        page = next((item for item in pages if item.number == row["page"]), None)
        if page is None:
            continue
        line_words = [word for word in page.words
                      if word.get("block", 0) == row.get("block")
                      and word.get("paragraph", 1) == row.get("paragraph")
                      and word.get("line", 0) == row.get("line")]
        cursor = 0
        selected_start = selected_end = None
        for word in line_words:
            word_start, word_end = cursor, cursor + len(word["text"])
            covered = any(mask["page"] == row["page"]
                          and word["x"] < mask["x"] + mask["width"]
                          and word["x"] + word["w"] > mask["x"]
                          and word["y"] < mask["y"] + mask["height"]
                          and word["y"] + word["h"] > mask["y"]
                          for mask in masks)
            if covered:
                if selected_start is None:
                    selected_start = word_start
                selected_end = word_end
            elif selected_start is not None:
                manual_spans.append((offsets[row["id"]] + selected_start,
                                     offsets[row["id"]] + selected_end,
                                     "MASKED_TEXT", row["text"][selected_start:selected_end]))
                selected_start = selected_end = None
            cursor = word_end + 1
        if selected_start is not None:
            manual_spans.append((offsets[row["id"]] + selected_start,
                                 offsets[row["id"]] + selected_end,
                                 "MASKED_TEXT", row["text"][selected_start:selected_end]))
    detected_spans = _sensitive_spans(raw_text)
    local_spans, approved_transmissions = list(manual_spans), []
    spans_by_row: dict[str, list[tuple[int, int, str, str]]] = {row["id"]: [] for row in rows}
    for span in detected_spans:
        overlapping = []
        for row in rows:
            start = offsets[row["id"]]
            end = start + len(row["text"])
            if span[0] < end and span[1] > start:
                overlapping.append(row)
                spans_by_row[row["id"]].append(span)
                row["detection_status"] = "sensitive"
                row["category"] = ", ".join(dict.fromkeys(filter(None, [row["category"], span[2]])))
        # In the simplified review flow, no OCR line is an explicit transmission
        # choice. A span discovered only after joining OCR fragments must therefore
        # stay local, even if each fragment was defaulted to "include". The older
        # detailed review flow may transmit it only after its explicit confirmation.
        if any(span[0] < manual_end and span[1] > manual_start
               for manual_start, manual_end, _, _ in manual_spans):
            local_spans.append(span)
        elif all(row["transmission"] == "include" for row in overlapping):
            if ocr_review is None:
                local_spans.append(span)
            elif not include_confirmed:
                raise ValueError("Explicit confirmation is required before sending selected sensitive OCR text to Gemini.")
            else:
                approved_transmissions.append(span[3])
        else:
            local_spans.append(span)
    for row in rows:
        start = offsets[row["id"]]
        row_spans = spans_by_row[row["id"]]
        manually_marked = row["manual_sensitive"] or row["classification"] == "sensitive"
        if manually_marked and not row_spans and row["text"] and not row.get("manual_mask_id"):
            manual_span = (start, start + len(row["text"]), "MANUAL", row["text"])
            row_spans = [manual_span]
            row["detection_status"] = "sensitive"
            row["category"] = row["category"] or "Manually marked"
            if row["transmission"] == "include" and ocr_review is not None:
                if not include_confirmed:
                    raise ValueError("Explicit confirmation is required before sending selected sensitive OCR text to Gemini.")
                approved_transmissions.append(row["text"])
            else:
                local_spans.append(manual_span)
        elif row_spans and row["transmission"] == "local":
            for span in row_spans:
                if span not in local_spans:
                    local_spans.append(span)
        if row_spans:
            row["detection_status"] = "sensitive"
        if manually_marked:
            row["detection_status"] = "sensitive"
    # A sensitive span can overlap a generic manual-mask token. Merge those spans
    # locally so token replacement never leaves part of the covered OCR value plain.
    merged_spans: list[tuple[int, int, str, str]] = []
    for span_start, span_end, kind, _ in sorted(local_spans, key=lambda item: (item[0], item[1])):
        if merged_spans and span_start < merged_spans[-1][1]:
            old_start, old_end, old_kind, _ = merged_spans[-1]
            merged_kind = "MASKED_TEXT" if "MASKED_TEXT" in {old_kind, kind} else old_kind
            merged_spans[-1] = (old_start, max(old_end, span_end), merged_kind,
                                raw_text[old_start:max(old_end, span_end)])
        else:
            merged_spans.append((span_start, span_end, kind, raw_text[span_start:span_end]))
    safe_text, mapping, findings, token_spans = _tokenize_spans(raw_text, merged_spans)
    # Detect every value in source OCR independently of the transmission choice so
    # its source image remains redacted even when text is explicitly approved.
    source_text = "\n".join(page.text for page in pages)
    source_visual_text, visual_mapping, _ = detect_and_tokenize(
        source_text, _manual_mask_spans(pages, masks))
    disabled = set(disabled_automatic_mask_ids or [])
    row_mask_ids = {f"auto-{row['id']}" for row in rows
                    if row["detection_status"] == "sensitive" and row["x"] is not None}
    source_disabled = [mask_id for mask_id in disabled
                       if not isinstance(mask_id, str) or not mask_id.startswith("auto-ocr-")]
    row_overrides = {item.get("id"): item for item in (automatic_mask_overrides or [])
                     if isinstance(item, dict) and str(item.get("id", "")).startswith("auto-ocr-")}
    source_overrides = [item for item in (automatic_mask_overrides or [])
                        if not isinstance(item, dict) or not str(item.get("id", "")).startswith("auto-ocr-")]
    automatic_masks, selected_auto_masks = _prepare_auto_masks(
        pages, visual_mapping, source_disabled, source_overrides)
    row_masks = []
    for row in rows:
        if row["detection_status"] != "sensitive" or row["x"] is None:
            continue
        mask_id = f"auto-{row['id']}"
        row["mask_id"] = mask_id
        coords = {"page": row["page"], "x": row["x"], "y": row["y"],
                  "width": row["width"], "height": row["height"]}
        if mask_id in row_overrides:
            override = _normalize_manual_masks(pages, [{**row_overrides[mask_id], "page": row["page"]}])[0]
            coords = {key: override[key] for key in ("page", "x", "y", "width", "height")}
        row_masks.append({"id": mask_id, **coords, "kind": "automatic",
                          "category": row["category"] or "Sensitive", "required": True,
                          "enabled": mask_id not in disabled})
    automatic_masks.extend(row_masks)
    selected_auto_masks.extend(mask for mask in row_masks if mask["enabled"])
    previews, visual_verified = _redact_visuals(pages, visual_mapping, selected_auto_masks, masks)
    if unmasked_manual_rows:
        visual_verified = False
    review_warnings = [f"Sensitive manual OCR item {row_id} needs an associated page mask." for row_id in unmasked_manual_rows]
    disabled_required_rows = [mask["id"] for mask in row_masks if mask["required"] and not mask["enabled"]]
    if disabled_required_rows:
        visual_verified = False
        review_warnings.extend(f"Required OCR mask {mask_id} is disabled." for mask_id in disabled_required_rows)
    if not visual_verified and not review_warnings:
        review_warnings.append("Sensitive OCR content is not fully covered by enabled image masks. Restore the suggested mask or draw a replacement mask over the indicated text.")
    preview_paths = [str(TEMP_DIR / Path(url).parts[-2] / Path(url).name) for url in previews]
    preview_hashes = [_file_sha256(Path(item)) for item in preview_paths]
    payload = build_gemini_payload(safe_text, output_language)
    page_numbers = [page.number for page in pages]
    privacy_verified = verify_privacy(payload, mapping, visual_verified,
                                      preview_paths, preview_hashes, masks, page_numbers,
                                      approved_transmissions)
    token_by_span = {(start, end): token for start, end, token in token_spans}
    for row in rows:
        start, end = offsets[row["id"]], offsets[row["id"]] + len(row["text"])
        row["tokens"] = [token for (span_start, span_end), token in token_by_span.items()
                         if span_start >= start and span_end <= end]
        row["token"] = ", ".join(row["tokens"]) or None
        row["will_send_to_ai"] = row["transmission"] == "include"
    return {"safe_ocr_text": safe_text, "sanitized_payload": payload,
            "redacted_previews": previews, "findings": findings,
            "ocr_items": rows, "approved_transmissions": approved_transmissions,
            "requires_transmission_confirmation": bool(approved_transmissions),
            "review_warnings": review_warnings,
            "output_language": output_language,
            "page_count": len(pages), "manual_masks": masks, "automatic_masks": automatic_masks,
            "ocr_diagnostics": [page.ocr_diagnostics for page in pages if page.ocr_diagnostics],
            "disabled_automatic_mask_ids": disabled_automatic_mask_ids or [],
            "automatic_mask_overrides": automatic_mask_overrides or [],
            "keep_pages": page_numbers,
            "page_dimensions": [{"page": page.number, "width": page.image.shape[1],
                                 "height": page.image.shape[0]} for page in pages],
            "privacy_verified": privacy_verified,
            "_ocr_cache": page_ocr_cache,
            "_token_map": mapping, "_visual_verified": visual_verified,
            "_approved_transmissions": approved_transmissions,
            "_page_numbers": page_numbers,
            "_sanitized_image_paths": preview_paths, "_sanitized_image_hashes": preview_hashes}


def extract_document(payload: str, mapping: dict[str, str], visual_verified: bool,
                     image_paths: list[str] | None = None, image_hashes: list[str] | None = None,
                     manual_masks: list[dict[str, int]] | None = None,
                     page_numbers: list[int] | None = None,
                     approved_transmissions: list[str] | None = None,
                     human_approved: bool = False) -> tuple[dict[str, Any], str]:
    """Call Gemini after explicit approval, rechecking privacy unless overridden."""
    if not image_paths:
        raise ValueError("Privacy verification failed. The document was not sent to external AI.")
    if not human_approved and not verify_privacy(payload, mapping, visual_verified, image_paths, image_hashes,
                                                 manual_masks, page_numbers, approved_transmissions):
        raise ValueError("Privacy verification failed. The document was not sent to external AI.")
    result, mode = call_gemini(payload, mapping, visual_verified, image_paths, image_hashes,
                               manual_masks, page_numbers, approved_transmissions, human_approved)
    if not result or not isinstance(result, dict):
        raise ValueError("AI returned an empty or invalid extraction")
    return result, mode


def export_excel(review: dict[str, Any], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Extracted data"
    result = restore_tokens(review["result"], review.get("_token_map", {}))
    fields = result.get("fields", {})
    if isinstance(fields, list):  # Compatibility with the original demo response.
        fields = {str(item.get("label", f"field_{index + 1}")): item.get("value", "")
                  for index, item in enumerate(fields) if isinstance(item, dict)}
    if not isinstance(fields, dict):
        fields = {"fields": fields}
    selected_fields = review.get("export_fields")
    if selected_fields is not None:
        allowed = set(selected_fields)
        fields = {key: value for key, value in fields.items() if key in allowed}
    arabic = review.get("output_language") == "ar"
    headers = (["نوع المستند", "الملخص", *[str(key) for key in fields]] if arabic
               else ["Document Type", "Summary", *[str(key) for key in fields]])
    sheet.append(headers)
    def cell_value(value: Any) -> Any:
        return json.dumps(value, ensure_ascii=False, indent=2) if isinstance(value, (dict, list)) else value
    sheet.append([cell_value(result.get("document_type", "")), cell_value(result.get("summary", "")),
                  *[cell_value(value) for value in fields.values()]])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1B5945")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        values = [len(str(cell.value or "")) for cell in column]
        sheet.column_dimensions[column[0].column_letter].width = min(max(max(values, default=12) + 2, 14), 48)
    findings = workbook.create_sheet("Locally redacted values" if not arabic else "القيم المحجوبة محليًا")
    findings.append(["الرمز", "الفئة", "النص الأصلي"] if arabic
                    else ["Token", "Category", "Original OCR text"])
    mapping = review.get("_token_map", {})
    categories = {item.get("token"): item.get("category", "Sensitive value")
                  for item in review.get("findings", [])}
    for token, original in mapping.items():
        findings.append([token, categories.get(token, "Sensitive value"), original])
    for cell in findings[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1B5945")
    findings.freeze_panes = "A2"
    findings.auto_filter.ref = findings.dimensions
    path.parent.mkdir(exist_ok=True)
    workbook.save(path)


def cleanup(path: Path | None = None) -> None:
    if path and path.exists():
        path.unlink(missing_ok=True)
    # Keep recent redacted previews accessible during the session; remove stale temp directories.
    if TEMP_DIR.exists():
        import time
        cutoff = time.time() - 24 * 60 * 60
        for child in TEMP_DIR.iterdir():
            try:
                if child.stat().st_mtime < cutoff:
                    shutil.rmtree(child, ignore_errors=True) if child.is_dir() else child.unlink(missing_ok=True)
            except OSError:
                pass
