"""Synthetic checks for the two human approval gates and privacy boundary."""
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
import uuid
import sys
import os
from types import SimpleNamespace
from contextlib import nullcontext
from unittest.mock import Mock, patch

import cv2
import numpy as np
from openpyxl import load_workbook
from PIL import Image

from backend import pipeline
from backend.pipeline import Page
from app import create_app, gemini_error_details, jobs

SYNTHETIC_OCR = (
    "Name: Test Person Employee ID: 9999999999 Phone: 0500000000 "
    "Email: test@example.com Address: 123 Test Road, Riyadh Department: Human Resources"
)
CREATED_TEST_IMAGES = []


def local_result():
    safe_text, mapping, findings = pipeline.detect_and_tokenize(SYNTHETIC_OCR)
    pipeline.TEMP_DIR.mkdir(exist_ok=True)
    image_path = pipeline.TEMP_DIR / f"test-redacted-{uuid.uuid4().hex}.png"
    cv2.imwrite(str(image_path), np.zeros((4, 4, 3), dtype=np.uint8))
    CREATED_TEST_IMAGES.append(image_path)
    return {
        "safe_ocr_text": safe_text,
        "sanitized_payload": pipeline.build_gemini_payload(safe_text),
        "redacted_previews": ["/api/preview/fake/redacted-1.png"],
        "findings": findings,
        "page_count": 1,
        "page_dimensions": [{"page": 0, "width": 4, "height": 4}],
        "keep_pages": [0],
        "manual_masks": [],
        "automatic_masks": [],
        "privacy_verified": True,
        "_page_numbers": [0],
        "_sanitized_image_paths": [str(image_path)],
        "_sanitized_image_hashes": [pipeline._file_sha256(image_path)],
        "_token_map": mapping,
        "_visual_verified": True,
    }


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.app = create_app()
        self.app.config["TESTING"] = True
        self.client = self.app.test_client()
        self.job_ids = []

    def tearDown(self):
        for job_id in self.job_ids:
            job = jobs.pop(job_id, {})
            source = job.get("_source_path")
            if source:
                Path(source).unlink(missing_ok=True)
            for image_path in job.get("_sanitized_image_paths", []):
                Path(image_path).unlink(missing_ok=True)
        for image_path in CREATED_TEST_IMAGES:
            Path(image_path).unlink(missing_ok=True)
        CREATED_TEST_IMAGES.clear()

    def test_gemini_request_errors_have_safe_actionable_messages(self):
        cases = [
            (429, "RESOURCE_EXHAUSTED", "gemini_quota", "quota was reached"),
            (403, "PERMISSION_DENIED", "gemini_auth", "permissions"),
            (404, "NOT_FOUND", "gemini_model", "configured model"),
            (503, "UNAVAILABLE", "gemini_unavailable", "temporarily unavailable"),
        ]
        for code, status, expected_code, phrase in cases:
            with self.subTest(code=code):
                error = pipeline.GeminiRequestError("ProviderError", code, status)
                actual_code, message = gemini_error_details(error)
                self.assertEqual(actual_code, expected_code)
                self.assertIn(phrase, message)
                self.assertNotIn("Test Person", message)

    def test_rechecks_reuse_server_local_ocr_without_running_tesseract_again(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.png"
            cv2.imwrite(str(source), np.full((32, 64, 3), 255, dtype=np.uint8))
            cached_words = [{"text": "Name", "x": 2, "y": 3, "w": 20, "h": 9,
                             "block": 1, "paragraph": 1, "line": 1, "word": 1,
                             "confidence": 99.0}]
            cache = [{"page": 0, "text": "Name", "words": cached_words,
                      "ocr_diagnostics": {"word_count": 1}}]
            with patch.object(pipeline.pytesseract, "get_languages", side_effect=AssertionError), \
                    patch.object(pipeline.pytesseract, "image_to_data", side_effect=AssertionError):
                pages = pipeline.load_document(source, cache)
            self.assertEqual(pages[0].text, "Name")
            self.assertEqual(pages[0].words, cached_words)
            self.assertEqual(pages[0].ocr_diagnostics["word_count"], 1)

    def upload_local(self, prepared=None, patch_pipeline=True):
        if patch_pipeline:
            context = patch("app.pipeline.prepare_document", return_value=prepared or local_result())
        else:
            context = nullcontext()
        with context:
            response = self.client.post(
                "/api/process", data={"file": (BytesIO(b"synthetic image"), "employee.png")},
                content_type="multipart/form-data")
        if response.is_json and response.json.get("job_id"):
            self.job_ids.append(response.json["job_id"])
        return response

    def test_contextual_english_and_arabic_names(self):
        labels = ["Name", "Full Name", "First Name", "Given Name", "Middle Name",
                  "Last Name", "Family Name", "Surname", "Recipient", "Recipient Name", "Holder",
                  "Holder Name", "Employee Name", "Customer Name", "Applicant",
                  "Applicant Name", "\u0627\u0644\u0627\u0633\u0645", "\u0627\u0644\u0627\u0633\u0645 \u0627\u0644\u0623\u0648\u0644",
                  "\u0627\u0644\u0627\u0633\u0645 \u0627\u0644\u0623\u0648\u0633\u0637",
                  "\u0627\u0644\u0627\u0633\u0645 \u0627\u0644\u0623\u062e\u064a\u0631",
                  "\u0627\u0633\u0645 \u0627\u0644\u0639\u0627\u0626\u0644\u0629", "\u0627\u0644\u0644\u0642\u0628",
                  "\u0627\u0633\u0645 \u0627\u0644\u0645\u0633\u062a\u0644\u0645",
                  "\u0627\u0633\u0645 \u0627\u0644\u0645\u0648\u0638\u0641",
                  "\u0627\u0633\u0645 \u062d\u0627\u0645\u0644",
                  "\u0627\u0644\u0645\u0633\u062a\u0644\u0645",
                  "\u0645\u0642\u062f\u0645 \u0627\u0644\u0637\u0644\u0628"]
        for label in labels:
            with self.subTest(label=label):
                safe, mapping, _ = pipeline.detect_and_tokenize(f"{label}: Test Person")
                self.assertEqual(mapping, {"[[PERSON_001]]": "Test Person"})
                self.assertEqual(safe, f"{label}: [[PERSON_001]]")

    def test_ocr_uses_both_languages_preserves_pdf_resolution_and_line_boundaries(self):
        words = ["Name", "الاسم", "Value"]
        ocr_data = {"text": words, "left": [10, 55, 10], "top": [12, 12, 35],
                    "width": [30, 36, 40], "height": [10, 10, 11],
                    "block_num": [1, 1, 2], "par_num": [1, 1, 1],
                    "line_num": [1, 1, 1], "word_num": [1, 2, 1],
                    "conf": [91, 88, 84]}
        class FakePage:
            def get_pixmap(self, **kwargs):
                self.render_args = kwargs
                return SimpleNamespace(width=2, height=2, n=3, samples=bytes([255] * 12))
        fake_page = FakePage()
        class FakeDocument:
            def __iter__(self): return iter([fake_page])
            def close(self): pass
        with tempfile.TemporaryDirectory() as directory, \
             patch.dict(os.environ, {"TESSERACT_LANG": "eng+ara"}), \
             patch.object(pipeline, "PDF_RENDER_DPI", 300), \
             patch.dict(sys.modules, {"fitz": SimpleNamespace(open=lambda path: FakeDocument())}), \
             patch.object(pipeline.pytesseract, "get_languages", return_value=["eng", "ara"]), \
             patch.object(pipeline.pytesseract, "image_to_data", return_value=ocr_data) as recognize:
            page = pipeline.load_document(Path(directory) / "sample.pdf")[0]
        self.assertEqual(fake_page.render_args["dpi"], pipeline.PDF_RENDER_DPI)
        self.assertGreaterEqual(fake_page.render_args["dpi"], 300)
        recognize.assert_called_once()
        self.assertEqual(recognize.call_args.kwargs["lang"], "eng+ara")
        self.assertEqual(recognize.call_args.kwargs["config"], "--oem 1 --psm 3")
        self.assertEqual(recognize.call_args.kwargs["output_type"], pipeline.pytesseract.Output.DICT)
        self.assertEqual(page.text, "Name الاسم\nValue")
        self.assertEqual([word["x"] for word in page.words], [10, 55, 10])
        self.assertEqual(page.ocr_diagnostics["line_count"], 2)

    def test_ocr_refuses_to_silently_drop_either_required_language_model(self):
        with patch.object(pipeline.pytesseract, "get_languages", return_value=["eng"]):
            with self.assertRaisesRegex(ValueError, "Arabic \\(ara\\)"):
                pipeline.load_document(Path("missing.png"))

    def test_requested_name_and_phone_employee_id_are_tokenized(self):
        text = "Name: Dhay M Alanazi Phone: 0501234567 Employee ID: 123456789"
        safe, mapping, findings = pipeline.detect_and_tokenize(text)
        self.assertEqual(len(findings), 3)
        self.assertEqual(mapping["[[PERSON_001]]"], "Dhay M Alanazi")
        self.assertEqual(mapping["[[PHONE_001]]"], "0501234567")
        self.assertEqual(mapping["[[EMPLOYEE_ID_001]]"], "123456789")
        for original in mapping.values():
            self.assertNotIn(original, safe)

    def test_local_ocr_text_is_tokenized_and_originals_stay_in_private_map(self):
        safe, mapping, findings = pipeline.detect_and_tokenize(SYNTHETIC_OCR)
        self.assertEqual(len(findings), 5)
        for private_value in ("Test Person", "9999999999", "0500000000", "test@example.com",
                              "123 Test Road, Riyadh"):
            self.assertNotIn(private_value, safe)
            self.assertIn(private_value, mapping.values())
        self.assertIn("[[PERSON_001]]", safe)
        self.assertIn("[[EMPLOYEE_ID_001]]", safe)
        self.assertEqual(pipeline.restore_tokens(safe, mapping), SYNTHETIC_OCR)

    def test_addresses_and_other_labeled_identifiers_are_masked(self):
        text = "Address: 123 Test Road, Riyadh National ID: 1234567890 Passport Number: AB1234567"
        safe, mapping, _ = pipeline.detect_and_tokenize(text)
        self.assertEqual(mapping["[[ADDRESS_001]]"], "123 Test Road, Riyadh")
        self.assertEqual(mapping["[[NATIONAL_ID_001]]"], "1234567890")
        self.assertEqual(mapping["[[PASSPORT_001]]"], "AB1234567")
        self.assertNotIn("123 Test Road", safe)
        self.assertNotIn("1234567890", safe)
        self.assertNotIn("AB1234567", safe)

    def test_payload_verification_checks_tokens_values_and_visual_flag(self):
        prepared = local_result()
        mapping = prepared["_token_map"]
        payload = prepared["sanitized_payload"]
        self.assertTrue(pipeline.verify_privacy(payload, mapping, True))
        for original in mapping.values():
            self.assertNotIn(original, payload)
        self.assertIn("[[PERSON_001]]", payload)
        self.assertFalse(pipeline.verify_privacy(payload + " Test Person", mapping, True))
        self.assertFalse(pipeline.verify_privacy(payload, mapping, False))
        with self.assertRaisesRegex(ValueError, "not sent to external AI"):
            pipeline.extract_document("Name: Test Person", mapping, True)

    def test_preparation_stops_at_local_boundary_and_ai_result_keeps_tokens(self):
        page = Page(np.full((60, 300, 3), 255, dtype=np.uint8), SYNTHETIC_OCR, [])
        captured = {}

        def fake_ai(payload, _mapping, _visual_verified, *_sanitized_images):
            captured["payload"] = payload
            return {"document_type": "employee_form", "summary": "Test form",
                    "fields": {"name": "[[PERSON_001]]", "phone": "[[PHONE_001]]"}}, "fake"

        with tempfile.TemporaryDirectory() as temporary_root:
            old_temp = pipeline.TEMP_DIR
            pipeline.TEMP_DIR = Path(temporary_root)
            def synthetic_preview(pages, _mapping, _automatic=None, _masks=None):
                directory = pipeline.TEMP_DIR / "fake"
                directory.mkdir()
                cv2.imwrite(str(directory / "redacted-1.png"), pages[0].image)
                return ["/api/preview/fake/redacted-1.png"], True
            try:
                with patch.object(pipeline, "load_document", return_value=[page]), \
                     patch.object(pipeline, "_redact_visuals", side_effect=synthetic_preview), \
                     patch.object(pipeline, "call_gemini", side_effect=fake_ai) as model:
                    prepared = pipeline.prepare_document(Path("synthetic.pdf"))
                    model.assert_not_called()
                    result, mode = pipeline.extract_document(prepared["sanitized_payload"],
                        prepared["_token_map"], prepared["_visual_verified"],
                        prepared["_sanitized_image_paths"], prepared["_sanitized_image_hashes"],
                        human_approved=True)
            finally:
                pipeline.TEMP_DIR = old_temp
        for original in prepared["_token_map"].values():
            self.assertNotIn(original, captured["payload"])
        self.assertIn("[[PERSON_001]]", captured["payload"])
        self.assertEqual(result["fields"]["name"], "[[PERSON_001]]")
        self.assertEqual(mode, "fake")

    def test_manual_masks_replace_ocr_text_and_verify_saved_image(self):
        words = [
            {"text": "Employee:", "x": 5, "y": 10, "w": 40, "h": 12},
            {"text": "Dhay", "x": 52, "y": 10, "w": 30, "h": 12},
            {"text": "M", "x": 88, "y": 10, "w": 10, "h": 12},
            {"text": "Alanazi", "x": 104, "y": 10, "w": 50, "h": 12},
            {"text": "Department:", "x": 170, "y": 10, "w": 60, "h": 12},
            {"text": "Ops", "x": 235, "y": 10, "w": 25, "h": 12},
        ]
        original_image = np.full((50, 280, 3), 255, dtype=np.uint8)
        page = Page(original_image.copy(), "Employee: Dhay M Alanazi Department: Ops", words)
        mask = {"page": 0, "x": 50, "y": 8, "width": 110, "height": 20}
        with tempfile.TemporaryDirectory() as temporary_root:
            old_temp = pipeline.TEMP_DIR
            pipeline.TEMP_DIR = Path(temporary_root)
            try:
                with patch.object(pipeline, "load_document", return_value=[page]):
                    prepared = pipeline.prepare_document(Path("synthetic.png"), [mask])
                token = "[[MASKED_TEXT_001]]"
                self.assertEqual(prepared["_token_map"][token], "Dhay M Alanazi")
                self.assertIn(token, prepared["safe_ocr_text"])
                self.assertNotIn("Dhay M Alanazi", prepared["safe_ocr_text"])
                self.assertTrue(pipeline.verify_privacy(prepared["sanitized_payload"],
                    prepared["_token_map"], prepared["_visual_verified"],
                    prepared["_sanitized_image_paths"], prepared["_sanitized_image_hashes"],
                    prepared["manual_masks"]))
                with patch.object(pipeline, "call_gemini", return_value=(
                        {"document_type": "form", "summary": "safe", "fields": {}}, "fake")) as model:
                    pipeline.extract_document(prepared["sanitized_payload"], prepared["_token_map"],
                        prepared["_visual_verified"], prepared["_sanitized_image_paths"],
                        prepared["_sanitized_image_hashes"], prepared["manual_masks"])
                self.assertEqual(model.call_args.args[3], prepared["_sanitized_image_paths"])
                self.assertNotIn("Dhay M Alanazi", model.call_args.args[0])
                redacted = cv2.imread(prepared["_sanitized_image_paths"][0])
                self.assertTrue(np.all(redacted[8:28, 50:160] == 0))
                self.assertTrue(np.all(page.image == 255))  # the source was not modified
                redacted[8, 50] = 255
                cv2.imwrite(prepared["_sanitized_image_paths"][0], redacted)
                self.assertFalse(pipeline.verify_privacy(prepared["sanitized_payload"],
                    prepared["_token_map"], prepared["_visual_verified"],
                    prepared["_sanitized_image_paths"], prepared["_sanitized_image_hashes"],
                    prepared["manual_masks"]))
            finally:
                pipeline.TEMP_DIR = old_temp

    def test_manual_mask_finds_exact_repeated_value_across_lines_and_pages(self):
        image = np.full((80, 260, 3), 255, dtype=np.uint8)
        first = Page(image.copy(), "Name\nAlice Smith", [
            {"text": "Name", "x": 5, "y": 5, "w": 32, "h": 12},
            {"text": "Alice", "x": 50, "y": 35, "w": 38, "h": 12},
            {"text": "Smith", "x": 92, "y": 35, "w": 40, "h": 12}], 0)
        second = Page(image.copy(), "Reviewed by Alice Smith", [
            {"text": "Reviewed", "x": 5, "y": 5, "w": 55, "h": 12},
            {"text": "by", "x": 64, "y": 5, "w": 12, "h": 12},
            {"text": "Alice", "x": 82, "y": 5, "w": 38, "h": 12},
            {"text": "Smith", "x": 124, "y": 5, "w": 40, "h": 12}], 1)
        masks = [{"id": "manual-name", "page": 0, "x": 48, "y": 33, "width": 88, "height": 16}]
        manual_spans = pipeline._manual_mask_spans([first, second], masks)
        self.assertEqual(manual_spans, [(5, 16, "PERSON_MANUAL", "Alice Smith")])
        source_text = "\n".join([first.text, second.text])
        safe, mapping, _ = pipeline.detect_and_tokenize(source_text, manual_spans)
        self.assertNotIn("Alice Smith", safe)
        self.assertIn("Alice Smith", mapping.values())
        repeated_masks = pipeline._automatic_masks([first, second], mapping)
        self.assertEqual([mask["page"] for mask in repeated_masks], [0, 1])

    def test_disabled_automatic_mask_can_be_replaced_by_manual_mask_and_page_order_is_preserved(self):
        image = np.full((40, 180, 3), 255, dtype=np.uint8)
        words = [{"text": "Employee:", "x": 5, "y": 10, "w": 40, "h": 12},
                 {"text": "Test", "x": 50, "y": 10, "w": 35, "h": 12},
                 {"text": "Person", "x": 90, "y": 10, "w": 45, "h": 12}]
        first = Page(image.copy(), "Employee: Test Person", words, 0)
        second = Page(np.full((20, 30, 3), 255, dtype=np.uint8), "Second page", [], 1)
        automatic_id = "auto-PERSON_MANUAL_001-p0-n1"
        with tempfile.TemporaryDirectory() as temporary_root:
            old_temp = pipeline.TEMP_DIR
            pipeline.TEMP_DIR = Path(temporary_root)
            try:
                with patch.object(pipeline, "load_document", return_value=[first, second]):
                    prepared = pipeline.prepare_document(Path("synthetic.pdf"),
                        [{"page": 0, "x": 48, "y": 8, "width": 90, "height": 17}],
                        keep_pages=[1, 0], disabled_automatic_mask_ids=[automatic_id])
                self.assertEqual(prepared["_page_numbers"], [1, 0])
                self.assertEqual([item["page"] for item in prepared["page_dimensions"]], [1, 0])
                self.assertFalse(next(item for item in prepared["automatic_masks"] if item["id"] == automatic_id)["enabled"])
                self.assertTrue(prepared["privacy_verified"])
                self.assertEqual(prepared["_token_map"]["[[MASKED_TEXT_001]]"], "Test Person")
            finally:
                pipeline.TEMP_DIR = old_temp

    def test_gemini_errors_distinguish_missing_key_and_bad_json(self):
        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "redacted.png"
            cv2.imwrite(str(image_path), np.zeros((20, 20, 3), dtype=np.uint8))
            image_hash = pipeline._file_sha256(image_path)
            with patch.dict("os.environ", {"DEMO_MODE": "false", "GEMINI_API_KEY": ""}):
                with self.assertRaises(pipeline.GeminiConfigurationError):
                    pipeline.call_gemini("safe", {}, True, [str(image_path)], [image_hash])
        parsed = pipeline._validate_json('```json\n{"document_type":"form","summary":"ok","fields":{}}\n```')
        self.assertEqual(parsed["document_type"], "form")
        for malformed in ("", "not json", '{"document_type":"form"}'):
            with self.subTest(response=malformed), self.assertRaises(pipeline.GeminiResponseError):
                pipeline._validate_json(malformed)

    def test_gemini_sdk_payload_excludes_local_token_mapping_and_original_masked_text(self):
        with tempfile.TemporaryDirectory() as directory:
            image_path = Path(directory) / "sanitized.png"
            cv2.imwrite(str(image_path), np.zeros((1200, 2400, 3), dtype=np.uint8))
            mapping = {"[[PERSON_MANUAL_001]]": "Original Secret Person"}
            payload = pipeline.build_gemini_payload("Name: [[PERSON_MANUAL_001]]", "ar")
            client = SimpleNamespace(models=SimpleNamespace(generate_content=Mock()))
            client.models.generate_content.return_value = SimpleNamespace(
                text='{"document_type":"Form","summary":"Name: [[PERSON_MANUAL_001]]","fields":{}}')
            with patch.dict("os.environ", {"GEMINI_API_KEY": "synthetic-test-key", "DEMO_MODE": "false"}), \
                    patch("google.genai.Client", return_value=client):
                pipeline.call_gemini(payload, mapping, True, [str(image_path)], [pipeline._file_sha256(image_path)])
                sent = client.models.generate_content.call_args.kwargs["contents"]
                config = client.models.generate_content.call_args.kwargs["config"]
                direct_payload = pipeline.build_gemini_payload("Name: Original Secret Person", "en")
                pipeline.call_gemini(direct_payload, {}, True, [str(image_path)],
                    [pipeline._file_sha256(image_path)], approved_transmissions=["Original Secret Person"])
            self.assertEqual(sent[0], payload)
            self.assertEqual(config["thinking_config"]["thinking_level"], "low")
            self.assertNotIn("Original Secret Person", str(sent))
            self.assertNotIn(str(mapping), str(sent))
            self.assertIn("[[PERSON_MANUAL_001]]", sent[0])
            self.assertEqual(sent[1].inline_data.mime_type, "image/jpeg")
            with Image.open(BytesIO(sent[1].inline_data.data)) as transmitted_image:
                self.assertEqual(transmitted_image.size, (1400, 700))
            direct_sent = client.models.generate_content.call_args.kwargs["contents"]
            self.assertIn("Original Secret Person", direct_sent[0])
            self.assertNotIn(str(mapping), str(direct_sent))

    def test_redacted_preview_verifies_and_blacks_out_sensitive_boxes(self):
        image = np.full((50, 140, 3), 255, dtype=np.uint8)
        words = [{"text": "Name:", "x": 5, "y": 10, "w": 28, "h": 12},
                 {"text": "Test", "x": 40, "y": 10, "w": 25, "h": 12},
                 {"text": "Person", "x": 68, "y": 10, "w": 35, "h": 12}]
        with tempfile.TemporaryDirectory(dir=pipeline.TEMP_DIR) as temporary_root:
            old_temp = pipeline.TEMP_DIR
            try:
                pipeline.TEMP_DIR = Path(temporary_root) / "generated"
                page = Page(image, "Name: Test Person", words)
                mappings = {"[[PERSON_001]]": "Test Person"}
                previews, verified = pipeline._redact_visuals(
                    [page], mappings, pipeline._automatic_masks([page], mappings))
                self.assertTrue(verified)
                destination = pipeline.TEMP_DIR / previews[0].split("/")[-2] / previews[0].split("/")[-1]
                redacted = cv2.imread(str(destination))
                self.assertTrue(np.all(redacted[10:22, 40:103] == 0))
            finally:
                pipeline.TEMP_DIR = old_temp

    def test_excel_has_dynamic_columns_nested_values_and_restores_tokens(self):
        review = {"result": {"document_type": "employee_form", "summary": "Test form",
                             "fields": {"name": "[[PERSON_001]]", "employee_id": "[[EMPLOYEE_ID_001]]",
                                        "phone": "[[PHONE_001]]", "department": "Human Resources",
                                        "metadata": {"source": "scan", "tags": ["approved", "2026"]}}},
                  "_token_map": {"[[PERSON_001]]": "Test Person",
                                 "[[EMPLOYEE_ID_001]]": "9999999999",
                                 "[[PHONE_001]]": "0500000000"}}
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "review.xlsx"
            pipeline.export_excel(review, destination)
            workbook = load_workbook(destination, data_only=True)
            sheet = workbook["Extracted data"]
            headers = [cell.value for cell in sheet[1]]
            row = [cell.value for cell in sheet[2]]
            values = dict(zip(headers, row))
            self.assertEqual(headers[:2], ["Document Type", "Summary"])
            self.assertEqual(headers[2:], ["name", "employee_id", "phone", "department", "metadata"])
            self.assertEqual(values["name"], "Test Person")
            self.assertEqual(values["employee_id"], "9999999999")
            self.assertEqual(values["phone"], "0500000000")
            self.assertIn('"source": "scan"', values["metadata"])
            self.assertFalse(any(isinstance(value, str) and value.startswith('{"document_type"') for value in row))
            workbook.close()

    def test_api_requires_pre_ai_approval_then_final_approval_before_export(self):
        response = self.upload_local()
        self.assertEqual(response.status_code, 200)
        job_id = response.json["job_id"]
        self.assertEqual(response.json["status"], "awaiting_ai_approval")
        self.assertEqual(response.json["metadata"]["privacy_status"], "PASSED")
        self.assertNotIn("_token_map", response.json)
        self.assertEqual(self.client.get(f"/api/export/{job_id}").status_code, 409)
        with patch("app.pipeline.extract_document") as premature_ai:
            self.assertEqual(self.client.post(f"/api/ai/{job_id}", json={}).status_code, 409)
            premature_ai.assert_not_called()
        with patch("app.pipeline.extract_document", return_value=(
                {"document_type": "employee_form", "summary": "Test form",
                 "fields": {"name": "[[PERSON_001]]", "department": "Human Resources"}}, "fake")) as ai:
            ai_response = self.client.post(f"/api/ai/{job_id}", json={"approved": True})
        self.assertEqual(ai_response.status_code, 200)
        self.assertEqual(ai.call_count, 1)
        self.assertIn("[[PERSON_001]]", ai.call_args.args[0])
        self.assertNotIn("Test Person", ai.call_args.args[0])
        self.assertEqual(ai_response.json["status"], "awaiting_final_approval")
        self.assertIn("[[PERSON_001]]", ai_response.json["ai_result"]["fields"]["name"])
        self.assertEqual(self.client.get(f"/api/export/{job_id}").status_code, 409)
        edited = {"document_type": "employee_form", "summary": "Approved test form",
                  "fields": {"الاسم": "[[PERSON_001]]", "القسم": "People"}}
        approval = self.client.post(f"/api/review/{job_id}", json={
            "action": "approve", "result": edited, "output_language": "ar"})
        self.assertEqual(approval.status_code, 200)
        self.assertEqual(approval.json["status"], "approved")
        self.assertNotIn("_token_map", approval.json)
        export_response = self.client.get(f"/api/export/{job_id}")
        self.assertEqual(export_response.status_code, 200)
        output_path = pipeline.OUTPUT_DIR / f"{job_id}.xlsx"
        workbook = load_workbook(output_path, data_only=True)
        sheet = workbook["البيانات المستخرجة"]
        exported = dict(zip([cell.value for cell in sheet[1]], [cell.value for cell in sheet[2]]))
        self.assertEqual(exported["الاسم"], "Test Person")
        self.assertEqual(exported["القسم"], "People")
        workbook.close()
        export_response.close()
        word_response = self.client.get(f"/api/export/{job_id}?format=word")
        self.assertEqual(word_response.status_code, 200)
        self.assertTrue(word_response.headers["Content-Disposition"].endswith("digitized-document.docx"))
        from docx import Document
        word_path = pipeline.OUTPUT_DIR / f"{job_id}.docx"
        word = Document(word_path)
        word_text = "\n".join(cell.text for row in word.tables[0].rows for cell in row.cells)
        self.assertIn("Test Person", word_text)
        self.assertIn("People", word_text)
        word_response.close()
        output_path.unlink(missing_ok=True)
        word_path.unlink(missing_ok=True)

    def test_reject_retains_local_state_and_retry_requires_local_review_again(self):
        uploaded = self.upload_local()
        job_id = uploaded.json["job_id"]
        original_path = Path(jobs[job_id]["_source_path"])
        self.assertTrue(original_path.exists())
        rejected = self.client.post(f"/api/review/{job_id}", json={"action": "reject"})
        self.assertEqual(rejected.json["status"], "rejected")
        self.assertTrue(original_path.exists())
        with patch("app.pipeline.extract_document") as ai:
            self.assertEqual(self.client.post(f"/api/ai/{job_id}", json={}).status_code, 409)
            ai.assert_not_called()
        retry = self.client.post(f"/api/review/{job_id}", json={"action": "retry"})
        self.assertEqual(retry.json["status"], "awaiting_ai_approval")
        self.assertTrue(retry.json["safe_ocr_text"])

    def test_manual_mask_api_accepts_only_coordinates_and_rebuilds_local_privacy_state(self):
        uploaded = self.upload_local()
        job_id = uploaded.json["job_id"]
        coordinates = [{"id": "manual-test", "page": 0, "x": 2, "y": 3, "width": 12, "height": 14}]
        prepared = local_result()
        prepared["manual_masks"] = coordinates
        prepared["page_dimensions"] = [{"page": 0, "width": 300, "height": 60}]
        prepared["_visual_verified"] = True
        prepared["privacy_verified"] = True
        with patch("app.pipeline.prepare_document", return_value=prepared) as local_processing:
            response = self.client.post(f"/api/manual-masks/{job_id}", json={"manual_masks": coordinates})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(local_processing.call_args.args[1], coordinates)
        self.assertEqual(response.json["manual_masks"][0]["id"], "manual-test")
        self.assertEqual(response.json["metadata"]["manual_mask_count"], 1)
        self.assertNotIn("_token_map", response.json)

    def test_manual_ocr_token_is_local_and_masked_text_is_not_in_sanitized_payload(self):
        image = np.full((40, 180, 3), 255, dtype=np.uint8)
        words = [{"text": "Employee:", "x": 5, "y": 10, "w": 40, "h": 12},
                 {"text": "Test", "x": 50, "y": 10, "w": 35, "h": 12},
                 {"text": "Person", "x": 90, "y": 10, "w": 45, "h": 12}]
        page = Page(image, "Employee: Test Person", words, 0)
        with tempfile.TemporaryDirectory() as directory:
            old_temp = pipeline.TEMP_DIR
            pipeline.TEMP_DIR = Path(directory)
            try:
                with patch.object(pipeline, "load_document", return_value=[page]):
                    prepared = pipeline.prepare_document(Path("synthetic.png"), [
                        {"id": "manual-private", "page": 0, "x": 48, "y": 8, "width": 90, "height": 17}])
                token = "[[MASKED_TEXT_001]]"
                self.assertIn(token, prepared["safe_ocr_text"])
                self.assertNotIn("Test Person", prepared["safe_ocr_text"])
                self.assertNotIn("Test Person", prepared["sanitized_payload"])
                self.assertEqual(prepared["_token_map"][token], "Test Person")
                self.assertNotIn("_token_map", {key: value for key, value in prepared.items() if not key.startswith("_")})
                export_path = Path(directory) / "manual-value.xlsx"
                pipeline.export_excel({"result": {"document_type": "Employee form", "summary": "Review",
                    "fields": {"department": "Operations"}}, "findings": prepared["findings"],
                    "_token_map": prepared["_token_map"]}, export_path)
                workbook = load_workbook(export_path, data_only=True)
                local_values = workbook["Locally redacted values"]
                rows = list(local_values.iter_rows(min_row=2, values_only=True))
                self.assertIn((token, "MASKED_TEXT", "Test Person"), rows)
                self.assertEqual(workbook["Extracted data"]["C2"].value, "Operations")
                workbook.close()
            finally:
                pipeline.TEMP_DIR = old_temp

    def test_manual_mask_tokenizes_arbitrary_overlapping_ocr_and_restores_it_in_excel(self):
        image = np.full((45, 220, 3), 255, dtype=np.uint8)
        words = [
            {"text": "Public", "x": 5, "y": 12, "w": 38, "h": 12},
            {"text": "ProjectX", "x": 55, "y": 12, "w": 58, "h": 12},
            {"text": "Notes", "x": 130, "y": 12, "w": 42, "h": 12},
        ]
        page = Page(image, "Public ProjectX Notes", words, 0)
        mask = {"id": "manual-arbitrary", "page": 0, "x": 50, "y": 8, "width": 70, "height": 20}
        with tempfile.TemporaryDirectory() as directory:
            old_temp = pipeline.TEMP_DIR
            pipeline.TEMP_DIR = Path(directory)
            try:
                with patch.object(pipeline, "load_document", return_value=[page]):
                    prepared = pipeline.prepare_document(Path("synthetic.png"), [mask])
                token = "[[MASKED_TEXT_001]]"
                self.assertEqual(prepared["_token_map"][token], "ProjectX")
                self.assertIn(f"Public {token} Notes", prepared["sanitized_payload"])
                self.assertNotIn("ProjectX", prepared["sanitized_payload"])
                self.assertTrue(prepared["privacy_verified"])

                export_path = Path(directory) / "restored-mask.xlsx"
                pipeline.export_excel({"result": {"document_type": "Form", "summary": "Reviewed",
                    "fields": {"project": token}}, "_token_map": prepared["_token_map"],
                    "findings": prepared["findings"]}, export_path)
                workbook = load_workbook(export_path, data_only=True)
                self.assertEqual(workbook["Extracted data"]["C2"].value, "ProjectX")
                local_values = list(workbook["Locally redacted values"].iter_rows(min_row=2, values_only=True))
                self.assertIn((token, "MASKED_TEXT", "ProjectX"), local_values)
                workbook.close()
            finally:
                pipeline.TEMP_DIR = old_temp

    def test_ocr_item_false_positive_is_still_detected_and_sensitive_transmission_needs_confirmation(self):
        image = np.full((40, 180, 3), 255, dtype=np.uint8)
        words = [{"text": "Name:", "x": 5, "y": 10, "w": 36, "h": 12, "block": 1, "line": 1},
                 {"text": "Test", "x": 50, "y": 10, "w": 35, "h": 12, "block": 1, "line": 1},
                 {"text": "Person", "x": 90, "y": 10, "w": 45, "h": 12, "block": 1, "line": 1}]
        page = Page(image, "Name: Test Person", words, 0)
        with tempfile.TemporaryDirectory() as directory:
            old_temp = pipeline.TEMP_DIR
            pipeline.TEMP_DIR = Path(directory)
            try:
                with patch.object(pipeline, "load_document", return_value=[page]):
                    local = pipeline.prepare_document(Path("synthetic.png"), ocr_review=[{
                        "id": "ocr-p0-b1-l1", "text": "Name: Test Person",
                        "classification": "non-sensitive", "transmission": "local"}])
                    self.assertIn("[[PERSON_001]]", local["safe_ocr_text"])
                    self.assertEqual(local["ocr_items"][0]["detection_status"], "sensitive")
                    with self.assertRaisesRegex(ValueError, "Explicit confirmation"):
                        pipeline.prepare_document(Path("synthetic.png"), ocr_review=[{
                            "id": "ocr-p0-b1-l1", "text": "Name: Test Person",
                            "classification": "non-sensitive", "transmission": "include"}])
                    included = pipeline.prepare_document(Path("synthetic.png"), ocr_review=[{
                        "id": "ocr-p0-b1-l1", "text": "Name: Test Person",
                        "classification": "non-sensitive", "transmission": "include"}], include_confirmed=True)
                self.assertIn("Test Person", included["safe_ocr_text"])
                self.assertEqual(included["approved_transmissions"], ["Test Person"])
                self.assertEqual(included["_token_map"], {})
                self.assertTrue(included["privacy_verified"])
                image_path = Path(directory) / Path(included["redacted_previews"][0]).parts[-2] / Path(included["redacted_previews"][0]).name
                redacted = cv2.imread(str(image_path))
                self.assertTrue(np.all(redacted[10:22, 5:135] == 0))
            finally:
                pipeline.TEMP_DIR = old_temp

    def test_sensitive_name_spanning_ocr_fragments_stays_local_in_simplified_flow(self):
        image = np.full((60, 200, 3), 255, dtype=np.uint8)
        words = [
            {"text": "Name:", "x": 5, "y": 8, "w": 42, "h": 12, "block": 1, "line": 1},
            {"text": "Dhay", "x": 5, "y": 30, "w": 35, "h": 12, "block": 1, "line": 2},
            {"text": "Alanazi", "x": 46, "y": 30, "w": 58, "h": 12, "block": 1, "line": 2},
        ]
        page = Page(image, "Name:\nDhay Alanazi", words, 0)
        with tempfile.TemporaryDirectory() as directory:
            old_temp = pipeline.TEMP_DIR
            pipeline.TEMP_DIR = Path(directory)
            try:
                with patch.object(pipeline, "load_document", return_value=[page]):
                    prepared = pipeline.prepare_document(Path("synthetic.png"))
                self.assertIn("[[PERSON_001]]", prepared["safe_ocr_text"])
                self.assertNotIn("Dhay Alanazi", prepared["safe_ocr_text"])
                self.assertEqual(prepared["_token_map"]["[[PERSON_001]]"], "Dhay Alanazi")
                self.assertEqual(prepared["approved_transmissions"], [])
                self.assertTrue(prepared["privacy_verified"])
            finally:
                pipeline.TEMP_DIR = old_temp

    def test_export_fields_selection_excludes_unchecked_attributes(self):
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "selected.xlsx"
            pipeline.export_excel({"result": {"document_type": "Form", "summary": "Summary",
                "fields": {"name": "Example", "department": "Operations"}},
                "export_fields": ["department"]}, destination)
            workbook = load_workbook(destination, data_only=True)
            sheet = workbook["Extracted data"]
            self.assertEqual([cell.value for cell in sheet[1]], ["Document Type", "Summary", "department"])
            self.assertEqual(sheet[2][2].value, "Operations")
            workbook.close()

    def test_api_mask_update_can_delete_manual_and_automatic_masks(self):
        uploaded = self.upload_local()
        job_id = uploaded.json["job_id"]
        auto_id = "auto-PERSON_001-p0-n1"
        jobs[job_id]["manual_masks"] = [{"id": "manual-to-delete", "kind": "manual", "page": 0,
                                          "x": 1, "y": 1, "width": 2, "height": 2}]
        jobs[job_id]["automatic_masks"] = [{"id": auto_id, "kind": "automatic", "page": 0,
                                            "x": 1, "y": 1, "width": 2, "height": 2, "enabled": True}]
        prepared = local_result()
        prepared.update({"manual_masks": [], "automatic_masks": [{"id": auto_id, "page": 0,
            "x": 1, "y": 1, "width": 10, "height": 10, "kind": "automatic", "enabled": False}],
            "disabled_automatic_mask_ids": [auto_id], "page_dimensions": [{"page": 0, "width": 4, "height": 4}],
            "_visual_verified": True, "privacy_verified": True})
        with patch("app.pipeline.prepare_document", return_value=prepared) as rebuild:
            response = self.client.post(f"/api/manual-masks/{job_id}", json={
                "manual_masks": [], "disabled_automatic_mask_ids": [auto_id],
                "automatic_mask_overrides": [], "keep_pages": [0]})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(rebuild.call_args.args[1], [])
        self.assertEqual(rebuild.call_args.args[3], [auto_id])
        self.assertEqual(response.json["manual_masks"], [])
        self.assertFalse(response.json["automatic_masks"][0]["enabled"])

    def test_frontend_mask_modes_selection_deletion_and_error_retention_are_explicit(self):
        root = Path(__file__).parents[1]
        source = (root / "frontend" / "app.js").read_text(encoding="utf-8")
        html = (root / "frontend" / "index.html").read_text(encoding="utf-8")
        styles = (root / "frontend" / "style.css").read_text(encoding="utf-8")
        backend_source = (root / "backend" / "pipeline.py").read_text(encoding="utf-8")
        self.assertIn('let selectedMaskId = null;', source)
        self.assertIn('if (hit && interactionMode === "draw") return;', source)
        self.assertIn('selectedMaskId = hit?.id || null;', source)
        self.assertIn('mask.id === selectedMaskId', source)
        self.assertIn('editorState.manual_masks = editorState.manual_masks.filter(mask => mask.id !== selectedMaskId)', source)
        self.assertIn('editorState.disabled_automatic_mask_ids.push(selectedMaskId)', source)
        self.assertIn('Your editor changes are retained', source)
        self.assertIn('id="select-mask-button"', html)
        self.assertIn('id="draw-mask-button"', html)
        self.assertIn('classList.add("mask-overlay")', source)
        self.assertIn('scheduleEditorRecheck();', source)
        self.assertIn('recheckTimer = setTimeout', source)
        self.assertIn('interactionActive', source)
        self.assertIn('id="mask-spinner"', html)
        self.assertIn('"300"', backend_source)
        self.assertIn('page.get_pixmap(dpi=PDF_RENDER_DPI', backend_source)
        self.assertIn('.mask-preview{display:block;width:100%;max-width:100%;height:auto', styles)
        self.assertIn('setTimeout(() =>', source)
        self.assertNotIn('getContext("2d")', source)
        self.assertNotIn('id="ocr-items-body"', html)
        self.assertNotIn('id="outgoing-payload"', html)
        self.assertNotIn('function renderOcrItems(items)', source)
        self.assertIn("Advanced OCR Details", html)
        self.assertIn('id="send-ai-button" disabled><span data-i18n="Continue to AI">Continue to AI</span>', html)
        self.assertIn('id="export-excel-button"', html)
        self.assertIn('id="export-word-button"', html)
        self.assertIn('id="interface-language"', html)
        self.assertNotIn('<details class="advanced-ocr" open', html)

    def test_local_processing_failure_without_redacted_images_still_cannot_send(self):
        with patch("app.pipeline.prepare_document", side_effect=ValueError(
                "Privacy verification failed. The document was not sent to external AI.")):
            blocked = self.upload_local(patch_pipeline=False)
        self.assertEqual(blocked.status_code, 422)
        self.assertEqual(blocked.json["metadata"]["privacy_status"], "BLOCKED")
        self.assertEqual(blocked.json["error"], "Privacy verification failed. The document was not sent to external AI.")
        with patch("app.pipeline.extract_document") as ai:
            response = self.client.post(f"/api/ai/{blocked.json['job_id']}", json={"approved": True})
            self.assertEqual(response.status_code, 409)
            ai.assert_not_called()

    def test_continue_to_ai_is_human_override_when_local_verification_needs_attention(self):
        uploaded = self.upload_local()
        job_id = uploaded.json["job_id"]
        jobs[job_id]["privacy_verified"] = False
        jobs[job_id]["status"] = "privacy_blocked"
        jobs[job_id]["metadata"]["privacy_status"] = "BLOCKED"
        with patch("app.pipeline.verify_privacy") as recheck, \
                patch("app.pipeline.extract_document", return_value=({"fields": {}}, "fake")) as model:
            response = self.client.post(f"/api/ai/{job_id}", json={"approved": True})
        self.assertEqual(response.status_code, 200, response.json)
        self.assertTrue(model.call_args.kwargs["human_approved"])
        self.assertTrue(response.json["metadata"]["privacy_override"])
        recheck.assert_not_called()

    def test_continue_to_ai_approval_covers_sensitive_transmission_consent(self):
        uploaded = self.upload_local()
        job_id = uploaded.json["job_id"]
        jobs[job_id]["_approved_transmissions"] = ["Test Person"]
        with patch("app.pipeline.extract_document", return_value=(
                {"document_type": "form", "summary": "reviewed", "fields": {}}, "fake")) as model:
            response = self.client.post(f"/api/ai/{job_id}", json={
                "approved": True, "output_language": "en", "include_confirmed": False})
        self.assertEqual(response.status_code, 200, response.json)
        model.assert_called_once()
        self.assertTrue(model.call_args.kwargs["human_approved"])

    def test_public_api_hides_raw_ocr_rows_and_transmission_values(self):
        prepared = local_result()
        prepared["ocr_items"] = [{"text": "Test Person", "original_text": "Test Person"}]
        prepared["approved_transmissions"] = ["Test Person"]
        uploaded = self.upload_local(prepared)
        serialized = uploaded.get_data(as_text=True)
        self.assertNotIn("ocr_items", uploaded.json)
        self.assertNotIn("approved_transmissions", uploaded.json)
        self.assertNotIn("Test Person", serialized)

    def test_frontend_exposes_three_step_visual_review_and_export_workflow(self):
        response = self.client.get("/")
        html = response.get_data(as_text=True)
        response.close()
        self.assertIn("Visual Review", html)
        self.assertIn("Results &amp; Export", html)
        self.assertIn("Continue to AI", html)
        self.assertIn("Export to Excel", html)
        self.assertNotIn("Local OCR text review", html)
        self.assertIn("id=\"fields-body\"", html)


    def test_nested_token_restoration_and_arabic_excel_headers(self):
        mapping = {"[[PERSON_001]]": "الاسم التجريبي", "[[PHONE_001]]": "0501234567"}
        for value in ("الاسم التجريبي", "Example Name", 123, "", "[[PERSON_001]]"):
            expected = mapping.get(value, value) if isinstance(value, str) else value
            self.assertEqual(pipeline.restore_tokens(value, mapping), expected)
        nested = pipeline.restore_tokens({"result [[PERSON_001]]": ["[[PHONE_001]]", {"nested": "[[PERSON_001]]"}]}, mapping)
        self.assertEqual(nested, {"result الاسم التجريبي": ["0501234567", {"nested": "الاسم التجريبي"}]})
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "arabic.xlsx"
            pipeline.export_excel({"result": {"document_type": "نموذج", "summary": "ملخص",
                                               "fields": {"الاسم": "[[PERSON_001]]"}},
                                   "output_language": "ar", "_token_map": mapping}, destination)
            workbook = load_workbook(destination, data_only=True)
            self.assertEqual([cell.value for cell in workbook["البيانات المستخرجة"][1]],
                             ["نوع المستند", "الملخص", "الاسم"])
            self.assertEqual(workbook["البيانات المستخرجة"]["C2"].value, "الاسم التجريبي")
            workbook.close()

    def test_arabic_column_labels_do_not_translate_document_values(self):
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "arabic-labels.xlsx"
            pipeline.export_excel({"result": {
                "document_type": "Employee registration form",
                "summary": "Description: Regional Operations Manager",
                "fields": {"اسم الشركة": "Acme Example Ltd", "العنوان": "123 Main Road, Riyadh",
                           "رقم الهوية": "ID-007-0009", "التاريخ": "2026-09-14"}},
                "output_language": "ar", "_token_map": {}}, destination)
            workbook = load_workbook(destination, data_only=True)
            sheet = workbook["البيانات المستخرجة"]
            self.assertEqual([cell.value for cell in sheet[1]],
                ["نوع المستند", "الملخص", "اسم الشركة", "العنوان", "رقم الهوية", "التاريخ"])
            self.assertEqual([cell.value for cell in sheet[2]], [
                "Employee registration form", "Description: Regional Operations Manager",
                "Acme Example Ltd", "123 Main Road, Riyadh", "ID-007-0009", "2026-09-14"])
            workbook.close()

    def test_export_word_localizes_labels_and_restores_values_without_translation(self):
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "arabic.docx"
            pipeline.export_word({"result": {"document_type": "Employee form", "summary": "Operations summary",
                "fields": {"اسم الشركة": "Acme Ltd", "الهاتف": "00442071234567", "الاسم": "[[PERSON_001]]"}},
                "output_language": "ar", "_token_map": {"[[PERSON_001]]": "Jane Doe"}}, destination)
            from docx import Document
            document = Document(destination)
            rows = [[cell.text for cell in row.cells] for row in document.tables[0].rows]
            self.assertEqual(rows[0], ["الحقل", "القيمة"])
            self.assertIn(["نوع المستند", "Employee form"], rows)
            self.assertIn(["الملخص", "Operations summary"], rows)
            self.assertIn(["اسم الشركة", "Acme Ltd"], rows)
            self.assertIn(["الاسم", "Jane Doe"], rows)

    def test_unlabelled_names_and_international_phone_numbers_are_sensitive(self):
        for value in ("John Smith", "أحمد محمد", "+44 20 7123 4567", "0049 30 12345678", "+1 (415) 555-2671"):
            with self.subTest(value=value):
                safe, mapping, findings = pipeline.detect_and_tokenize(value)
                self.assertNotIn(value, safe)
                self.assertIn(value, mapping.values())
                self.assertTrue(findings)

    def test_coop_registration_heading_is_not_mistaken_for_a_person_name(self):
        for value in ("COOP REGISTRATION", "Coop Registration"):
            with self.subTest(value=value):
                safe, mapping, findings = pipeline.detect_and_tokenize(value)
                self.assertEqual(safe, value)
                self.assertEqual(mapping, {})
                self.assertEqual(findings, [])

    def test_dashboard_heading_is_not_mistaken_for_a_person_name(self):
        for value in ("الداش بورد الداخلي", "الداشبورد الداخلي", "internal dashboard"):
            with self.subTest(value=value):
                safe, mapping, findings = pipeline.detect_and_tokenize(value)
                self.assertEqual(safe, value)
                self.assertEqual(mapping, {})
                self.assertEqual(findings, [])

    def test_university_and_education_names_are_not_mistaken_for_people(self):
        for value in ("جامعة الملك عبدالعزيز", "الجامعة السعودية", "كلية علوم الحاسب",
                      "King Saud University", "Ministry of Education"):
            with self.subTest(value=value):
                safe, mapping, findings = pipeline.detect_and_tokenize(value)
                self.assertEqual(safe, value)
                self.assertEqual(mapping, {})
                self.assertEqual(findings, [])

    def test_short_office_and_heading_phrases_are_not_mistaken_for_people(self):
        for value in ("المكتب الداخلي", "مكتب التدريب", "The Office", "The Main Office",
                      "Dashboard Internal", "إدارة المكتب"):
            with self.subTest(value=value):
                safe, mapping, findings = pipeline.detect_and_tokenize(value)
                self.assertEqual(safe, value)
                self.assertEqual(mapping, {})
                self.assertEqual(findings, [])

    def test_company_representative_field_accepts_arabic_english_and_token_text(self):
        token_map = {"[[PERSON_001]]": "Test Person"}
        cases = [("الاسم التجريبي", "الاسم التجريبي"),
                 ("Example Person", "Example Person"),
                 ("[[PERSON_001]]", "Test Person")]
        with tempfile.TemporaryDirectory() as directory:
            for index, (submitted, expected) in enumerate(cases):
                with self.subTest(value=submitted):
                    destination = Path(directory) / f"field-{index}.xlsx"
                    pipeline.export_excel({"result": {"document_type": "form", "summary": "reviewed",
                        "fields": {"company_representative_name_page_2": submitted}},
                        "_token_map": token_map}, destination)
                    workbook = load_workbook(destination, data_only=True)
                    sheet = workbook["Extracted data"]
                    values = dict(zip([cell.value for cell in sheet[1]], [cell.value for cell in sheet[2]]))
                    self.assertEqual(values["company_representative_name_page_2"], expected)
                    workbook.close()

    def test_end_to_end_synthetic_upload_privacy_ai_review_and_excel(self):
        words = [{"text": "Name:", "x": 5, "y": 10, "w": 28, "h": 12},
                 {"text": "Test", "x": 40, "y": 10, "w": 25, "h": 12},
                 {"text": "Person", "x": 68, "y": 10, "w": 35, "h": 12}]
        original = np.full((40, 140, 3), 255, dtype=np.uint8)
        page = Page(original.copy(), "Name: Test Person", words, 0)
        with tempfile.TemporaryDirectory() as temporary_root:
            old_temp = pipeline.TEMP_DIR
            pipeline.TEMP_DIR = Path(temporary_root)
            output_path = None
            try:
                with patch.object(pipeline, "load_document", return_value=[page]):
                    uploaded = self.client.post("/api/process",
                        data={"file": (BytesIO(b"synthetic PNG"), "synthetic.png")},
                        content_type="multipart/form-data")
                self.assertEqual(uploaded.status_code, 200, uploaded.json)
                job_id = uploaded.json["job_id"]
                self.job_ids.append(job_id)
                self.assertEqual(uploaded.json["status"], "awaiting_ai_approval")
                self.assertNotIn("Test Person", uploaded.json["safe_ocr_text"])
                self.assertNotIn("_token_map", uploaded.json)
                with patch("app.pipeline.extract_document") as premature_ai:
                    self.assertEqual(self.client.post(f"/api/ai/{job_id}",
                        json={"output_language": "ar"}).status_code, 409)
                    premature_ai.assert_not_called()
                def fake_extract(payload, mapping, visual, images, hashes, masks, page_numbers,
                                 approved_transmissions=None, *, human_approved=False):
                    self.assertTrue(human_approved)
                    self.assertIn("Use Arabic only for field labels and other output labels.", payload)
                    self.assertIn("Do not translate, paraphrase", payload)
                    self.assertNotIn("Test Person", payload)
                    self.assertIn("[[PERSON_001]]", payload)
                    self.assertTrue(pipeline.verify_privacy(payload, mapping, visual, images, hashes,
                                                            masks, page_numbers, approved_transmissions))
                    sanitized = cv2.imread(images[0])
                    self.assertTrue(np.all(sanitized[10:22, 40:103] == 0))
                    return ({"document_type": "نموذج", "summary": "ملخص", "fields": {
                        "company_representative_name_page_2": "الاسم التجريبي",
                        "token_check": "[[PERSON_001]]"}}, "fake")
                with patch("app.pipeline.extract_document", side_effect=fake_extract) as model:
                    ai_response = self.client.post(f"/api/ai/{job_id}",
                        json={"approved": True, "output_language": "ar"})
                self.assertEqual(ai_response.status_code, 200, ai_response.json)
                self.assertEqual(ai_response.json["output_language"], "ar")
                self.assertEqual(model.call_count, 1)
                final_result = {"document_type": "نموذج", "summary": "ملخص",
                                "fields": {"company_representative_name_page_2": "الاسم التجريبي",
                                           "token_check": "[[PERSON_001]]"}}
                approval = self.client.post(f"/api/review/{job_id}",
                    json={"action": "approve", "result": final_result})
                self.assertEqual(approval.status_code, 200)
                exported = self.client.get(f"/api/export/{job_id}")
                self.assertEqual(exported.status_code, 200)
                output_path = pipeline.OUTPUT_DIR / f"{job_id}.xlsx"
                workbook = load_workbook(output_path, data_only=True)
                sheet = workbook["البيانات المستخرجة"]
                values = dict(zip([cell.value for cell in sheet[1]], [cell.value for cell in sheet[2]]))
                self.assertEqual(values["نوع المستند"], "نموذج")
                self.assertEqual(values["company_representative_name_page_2"], "الاسم التجريبي")
                self.assertEqual(values["token_check"], "Test Person")
                local_values = workbook["القيم المحجوبة محليًا"]
                self.assertEqual(local_values["C2"].value, "Test Person")
                workbook.close(); exported.close()
            finally:
                if output_path:
                    output_path.unlink(missing_ok=True)
                pipeline.TEMP_DIR = old_temp


if __name__ == "__main__":
    unittest.main()
